from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.db.models import AuditLog, ImportRun, PprEvent, PprNotification
from app.excel.importer import HEADER_ALIASES, cell, combine_datetime, get_header_map, parse_bool, parse_date, parse_time
from app.services.ppr_service import describe_changes, sync_start_notification
from app.services.statuses import (
    NOTIFICATION_STATUS_CANCELLED,
    NOTIFICATION_STATUS_PLANNED,
    NOTIFICATION_STATUS_SKIPPED,
    PPR_STATUS_ARCHIVED,
    PPR_STATUS_SCHEDULED,
)
from app.services.user_service import user_display

IMPORT_MODE_SAFE = "safe"
IMPORT_MODE_NEW_ONLY = "new_only"
IMPORT_MODE_FORCE = "force"
IMPORT_MODES = {IMPORT_MODE_SAFE, IMPORT_MODE_NEW_ONLY, IMPORT_MODE_FORCE}

IMPORT_STATUS_PREVIEW = "preview"
IMPORT_STATUS_COMPLETED = "completed"
IMPORT_STATUS_FAILED = "failed"

ACTION_CREATE = "create"
ACTION_UPDATE = "update"
ACTION_UNCHANGED = "unchanged"
ACTION_SKIP_MANUAL = "skip_manual"
ACTION_DUPLICATE = "duplicate"
ACTION_INVALID = "invalid"
ACTION_MISSING_FROM_EXCEL = "missing_from_excel"
ACTION_SOURCE_KEY_MIGRATION = "source_key_migration"
ACTION_AMBIGUOUS = "ambiguous"

KEY_METHOD_ID = "id"
KEY_METHOD_SOURCE = "source"
KEY_METHOD_FINGERPRINT = "fingerprint"

MATCH_METHOD_SOURCE_KEY = "source_key"
MATCH_METHOD_EXTERNAL_ID = "external_id"
MATCH_METHOD_LEGACY_ROW = "legacy_row"
MATCH_METHOD_FINGERPRINT = "fingerprint"
MATCH_METHOD_AMBIGUOUS = "ambiguous"
MATCH_METHOD_NONE = "none"

NOTIFICATION_CREATE = "create"
NOTIFICATION_UPDATE = "update"
NOTIFICATION_DISABLE_MISSING_DATE = "disable_missing_date"
NOTIFICATION_DISABLE_NOTIFY_FALSE = "disable_notify_false"
NOTIFICATION_SKIP = "skip"


class ImportPreviewError(ValueError):
    pass


class ImportPreviewNotFound(ImportPreviewError):
    pass


class ImportFileChanged(ImportPreviewError):
    pass


class ImportForceConfirmationRequired(ImportPreviewError):
    pass


class ImportRepeatedFile(ImportPreviewError):
    pass


@dataclass
class ParsedExcelRow:
    excel_row_number: int
    external_id: str
    source_key: str
    title: str
    project: str | None
    values: dict[str, Any]
    external_id_from_excel: str | None
    source_value: str | None
    key_method: str
    legacy_source_key: str | None


@dataclass
class MatchResult:
    events: list[PprEvent]
    method: str
    reason: str


def validate_mode(mode: str) -> str:
    normalized = (mode or IMPORT_MODE_SAFE).strip().lower()
    if normalized not in IMPORT_MODES:
        raise ImportPreviewError(f"Unsupported import mode: {mode}")
    return normalized


def file_sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def read_file_bytes(path: str | Path) -> tuple[str, bytes]:
    file_path = Path(path)
    return file_path.name, file_path.read_bytes()


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def normalized_string(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def fingerprint_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def stable_fingerprint_payload(values: dict[str, Any]) -> dict[str, str]:
    return {
        "title": fingerprint_value(values.get("title")),
        "project": fingerprint_value(values.get("project")),
        "notification_type": fingerprint_value(values.get("notification_type")),
        "activities": fingerprint_value(values.get("activities")),
        "responsible_setup": fingerprint_value(values.get("responsible_setup")),
        "responsible_report": fingerprint_value(values.get("responsible_report")),
        "source_link": fingerprint_value(values.get("source_link")),
    }


def build_fingerprint_source_key(values: dict[str, Any]) -> str:
    payload = stable_fingerprint_payload(values)
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return f"fingerprint:{hashlib.sha256(raw.encode('utf-8')).hexdigest()}"


def source_row_number(value: Any, excel_row_num: int) -> int:
    if value in (None, ""):
        return excel_row_num
    try:
        return int(value)
    except (TypeError, ValueError):
        return excel_row_num


def build_source_key(external_id_from_excel: str | None, source_value: str | None, values: dict[str, Any]) -> tuple[str, str]:
    if external_id_from_excel:
        return f"id:{external_id_from_excel}", KEY_METHOD_ID
    if source_value:
        return f"source:{source_value}", KEY_METHOD_SOURCE
    return build_fingerprint_source_key(values), KEY_METHOD_FINGERPRINT


def build_external_id(external_id_from_excel: str | None, excel_row_num: int) -> str:
    if external_id_from_excel:
        return external_id_from_excel
    return f"ROW-{excel_row_num}"


def automatic_external_id(source_key: str) -> str:
    digest = source_key.split(":", 1)[1] if ":" in source_key else hashlib.sha256(source_key.encode("utf-8")).hexdigest()
    return f"AUTO-{digest[:48]}"


def source_key_from_external_id(external_id: str) -> str | None:
    if external_id.startswith("MANUAL-"):
        return None
    if external_id.startswith("ROW-"):
        return f"row:{external_id[4:]}"
    return f"id:{external_id}"


def is_legacy_row_source_key(value: str | None) -> bool:
    return bool(value and value.startswith("row:"))


def event_fingerprint_source_key(event: PprEvent) -> str:
    values = {
        "title": event.title,
        "project": event.project,
        "notification_type": event.notification_type,
        "activities": event.activities,
        "responsible_setup": event.responsible_setup,
        "responsible_report": event.responsible_report,
        "source_link": event.source_link,
    }
    return build_fingerprint_source_key(values)


def is_manual_event(event: PprEvent) -> bool:
    return bool(event.external_id and event.external_id.startswith("MANUAL-"))


def row_has_values(row) -> bool:
    return any(item.value not in (None, "") for item in row)


def parse_excel_content(content: bytes, sheet_name: str = "ППР_для_бота") -> tuple[list[ParsedExcelRow], list[dict], int]:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    header_map = get_header_map(ws)
    if "title" not in header_map:
        raise ImportPreviewError("В Excel нет обязательной колонки 'Название ППР'")

    rows: list[ParsedExcelRow] = []
    invalid_rows: list[dict] = []
    total_rows = 0

    for excel_row_num, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if not row_has_values(row):
            continue
        total_rows += 1
        try:
            title = normalized_string(cell(row, header_map, "title"))
            if not title:
                raise ValueError("Название ППР не заполнено")

            external_id_from_excel = normalized_string(cell(row, header_map, "external_id"))
            source_value = normalized_string(cell(row, header_map, "source_row"))
            source_row = source_row_number(source_value, excel_row_num)
            external_id = build_external_id(external_id_from_excel, excel_row_num)
            event_date = parse_date(cell(row, header_map, "date"))
            start_time = parse_time(cell(row, header_map, "start_time"), default=time(8, 0, 0)) if event_date else None
            end_time = parse_time(cell(row, header_map, "end_time"), default=None)

            values = {
                "external_id": external_id,
                "source_key": None,
                "source_row": source_row,
                "date": event_date,
                "start_time": start_time,
                "end_time": end_time,
                "notification_type": normalized_string(cell(row, header_map, "notification_type")),
                "project": normalized_string(cell(row, header_map, "project")),
                "title": title,
                "activities": normalized_string(cell(row, header_map, "activities")),
                "responsible_setup": normalized_string(cell(row, header_map, "responsible_setup")),
                "responsible_report": normalized_string(cell(row, header_map, "responsible_report")),
                "source_link": normalized_string(cell(row, header_map, "source_link")),
                "notify_start": parse_bool(cell(row, header_map, "notify_start"), default=True),
                "notify_end": parse_bool(cell(row, header_map, "notify_end"), default=False),
                "is_active": parse_bool(cell(row, header_map, "is_active"), default=True),
                "comment": normalized_string(cell(row, header_map, "comment")),
            }
            source_key, key_method = build_source_key(external_id_from_excel, source_value, values)
            legacy_source_key = f"row:{source_row}"
            values["source_key"] = source_key
            if not external_id_from_excel:
                external_id = automatic_external_id(source_key)
                values["external_id"] = external_id
            rows.append(
                ParsedExcelRow(
                    excel_row_number=excel_row_num,
                    external_id=external_id,
                    source_key=source_key,
                    title=title,
                    project=values["project"],
                    values=values,
                    external_id_from_excel=external_id_from_excel,
                    source_value=source_value,
                    key_method=key_method,
                    legacy_source_key=legacy_source_key,
                )
            )
        except Exception as exc:
            invalid_rows.append(
                {
                    "excel_row_number": excel_row_num,
                    "external_id": None,
                    "source_key": None,
                    "title": normalized_string(cell(row, header_map, "title")),
                    "project": normalized_string(cell(row, header_map, "project")),
                    "current_date": None,
                    "current_time": None,
                    "new_date": None,
                    "new_time": None,
                    "action": ACTION_INVALID,
                    "key_method": MATCH_METHOD_NONE,
                    "match_method": MATCH_METHOD_NONE,
                    "confidence": "none",
                    "reason": str(exc),
                    "fields_changed": {},
                    "notification_action": NOTIFICATION_SKIP,
                    "notification_reason": "Строка невалидна",
                }
            )

    return rows, invalid_rows, total_rows


def find_matching_events(db: Session, row: ParsedExcelRow) -> MatchResult:
    source_key_events = db.query(PprEvent).filter(PprEvent.source_key == row.source_key).all() if row.source_key else []
    if len(source_key_events) > 1:
        return MatchResult(source_key_events, MATCH_METHOD_AMBIGUOUS, "source_key сопоставился с несколькими ППР")
    if len(source_key_events) == 1:
        return MatchResult(source_key_events, MATCH_METHOD_SOURCE_KEY, "Найдено по source_key")

    if row.external_id_from_excel:
        external_id_events = db.query(PprEvent).filter(PprEvent.external_id == row.external_id).all()
        if len(external_id_events) > 1:
            return MatchResult(external_id_events, MATCH_METHOD_AMBIGUOUS, "external_id сопоставился с несколькими ППР")
        if len(external_id_events) == 1:
            return MatchResult(external_id_events, MATCH_METHOD_EXTERNAL_ID, "Найдено по external_id; source_key можно обновить")

    legacy_filters = []
    if row.legacy_source_key:
        legacy_filters.append(PprEvent.source_key == row.legacy_source_key)
    legacy_external_id = f"ROW-{row.excel_row_number}"
    legacy_filters.append(PprEvent.external_id == legacy_external_id)
    legacy_events = db.query(PprEvent).filter(or_(*legacy_filters)).all() if legacy_filters else []
    legacy_by_id = {event.id: event for event in legacy_events if not is_manual_event(event)}
    if len(legacy_by_id) > 1:
        return MatchResult(list(legacy_by_id.values()), MATCH_METHOD_AMBIGUOUS, "legacy row сопоставился с несколькими ППР")
    if len(legacy_by_id) == 1:
        return MatchResult(list(legacy_by_id.values()), MATCH_METHOD_LEGACY_ROW, "Найдено по legacy row; source_key можно обновить")

    if row.key_method == KEY_METHOD_FINGERPRINT:
        fingerprint_matches: dict[int, PprEvent] = {}
        imported_events = db.query(PprEvent).filter(PprEvent.external_id.notlike("MANUAL-%")).all()
        for event in imported_events:
            if event.source_key == row.source_key:
                fingerprint_matches[event.id] = event
            elif event.source_key is None and event_fingerprint_source_key(event) == row.source_key:
                fingerprint_matches[event.id] = event
            elif is_legacy_row_source_key(event.source_key) and event_fingerprint_source_key(event) == row.source_key:
                fingerprint_matches[event.id] = event
        if len(fingerprint_matches) > 1:
            return MatchResult(list(fingerprint_matches.values()), MATCH_METHOD_AMBIGUOUS, "fingerprint сопоставился с несколькими ППР")
        if len(fingerprint_matches) == 1:
            return MatchResult(list(fingerprint_matches.values()), MATCH_METHOD_FINGERPRINT, "Найдено по уникальному fingerprint; source_key можно обновить")

    return MatchResult([], MATCH_METHOD_NONE, "Совпадений в БД нет")


def compare_event_values(event: PprEvent, values: dict[str, Any]) -> dict[str, tuple[Any, Any]]:
    changes: dict[str, tuple[Any, Any]] = {}
    for field in (
        "external_id",
        "source_key",
        "source_row",
        "date",
        "start_time",
        "end_time",
        "notification_type",
        "project",
        "title",
        "activities",
        "responsible_setup",
        "responsible_report",
        "source_link",
        "notify_start",
        "notify_end",
        "is_active",
        "comment",
    ):
        old_value = getattr(event, field)
        new_value = values[field]
        if old_value != new_value:
            changes[field] = (old_value, new_value)
    return changes


def fields_changed_json(changes: dict[str, tuple[Any, Any]]) -> dict[str, dict[str, Any]]:
    return {
        field: {"old": json_value(old_value), "new": json_value(new_value)}
        for field, (old_value, new_value) in changes.items()
    }


def first_start_notification(event: PprEvent | None) -> PprNotification | None:
    if event is None:
        return None
    for notif in event.notifications:
        if notif.type == "start":
            return notif
    return None


def notification_preview(event: PprEvent | None, values: dict[str, Any], action: str, mode: str) -> tuple[str, str]:
    if action in {ACTION_INVALID, ACTION_DUPLICATE, ACTION_SKIP_MANUAL, ACTION_MISSING_FROM_EXCEL, ACTION_SOURCE_KEY_MIGRATION, ACTION_AMBIGUOUS}:
        return NOTIFICATION_SKIP, "Уведомление не меняется для этой строки"
    if action == ACTION_UNCHANGED and mode == IMPORT_MODE_NEW_ONLY:
        return NOTIFICATION_SKIP, "new_only не меняет существующие ППР"

    scheduled_at = combine_datetime(values.get("date"), values.get("start_time"))
    notif = first_start_notification(event)

    if not values.get("date"):
        if notif:
            return NOTIFICATION_DISABLE_MISSING_DATE, "Дата отсутствует, существующее уведомление будет отключено"
        return NOTIFICATION_SKIP, "Дата отсутствует, уведомление не создается"

    if not values.get("is_active", True):
        if notif:
            return NOTIFICATION_UPDATE, "ППР неактивна, уведомление будет отменено"
        return NOTIFICATION_SKIP, "ППР неактивна, уведомление не создается"

    if not values.get("notify_start", True):
        if notif:
            return NOTIFICATION_DISABLE_NOTIFY_FALSE, "notify=false, существующее уведомление будет отключено"
        return NOTIFICATION_SKIP, "notify=false, уведомление не создается"

    if scheduled_at is None:
        return NOTIFICATION_SKIP, "Нет даты или времени выхода"

    auto_send_note = "auto_send выключен, потому что дата/время в прошлом" if scheduled_at <= datetime.now() else "auto_send будет включен"
    if notif is None:
        return NOTIFICATION_CREATE, f"Будет создано start-уведомление; {auto_send_note}"
    if notif.scheduled_at != scheduled_at or notif.status in {NOTIFICATION_STATUS_SKIPPED, NOTIFICATION_STATUS_CANCELLED}:
        return NOTIFICATION_UPDATE, f"Будет обновлено start-уведомление; {auto_send_note}"
    return NOTIFICATION_SKIP, "Уведомление не изменится"


def row_detail(
    row: ParsedExcelRow,
    event: PprEvent | None,
    action: str,
    reason: str,
    changes: dict[str, tuple[Any, Any]],
    mode: str,
    match_method: str | None = None,
    confidence: str | None = None,
) -> dict:
    notification_action, notification_reason = notification_preview(event, row.values, action, mode)
    return {
        "excel_row_number": row.excel_row_number,
        "external_id": row.external_id,
        "source_key": row.source_key,
        "title": row.title,
        "project": row.project,
        "current_date": event.date.isoformat() if event and event.date else None,
        "current_time": event.start_time.isoformat() if event and event.start_time else None,
        "new_date": row.values["date"].isoformat() if row.values["date"] else None,
        "new_time": row.values["start_time"].isoformat() if row.values["start_time"] else None,
        "action": action,
        "key_method": row.key_method,
        "match_method": match_method or MATCH_METHOD_NONE,
        "confidence": confidence or ("high" if (match_method or MATCH_METHOD_NONE) in {MATCH_METHOD_SOURCE_KEY, MATCH_METHOD_EXTERNAL_ID, MATCH_METHOD_FINGERPRINT} else "none"),
        "reason": reason,
        "fields_changed": fields_changed_json(changes),
        "notification_action": notification_action,
        "notification_reason": notification_reason,
        "new_values": {key: json_value(value) for key, value in row.values.items()},
        "ppr_event_id": event.id if event else None,
    }


def missing_from_excel_detail(event: PprEvent) -> dict:
    return {
        "excel_row_number": None,
        "external_id": event.external_id,
        "source_key": event.source_key,
        "title": event.title,
        "project": event.project,
        "current_date": event.date.isoformat() if event.date else None,
        "current_time": event.start_time.isoformat() if event.start_time else None,
        "new_date": None,
        "new_time": None,
        "action": ACTION_MISSING_FROM_EXCEL,
        "key_method": MATCH_METHOD_NONE,
        "match_method": MATCH_METHOD_NONE,
        "confidence": "none",
        "reason": "ППР есть в БД, но отсутствует в текущем Excel. Автоматически не архивируется.",
        "fields_changed": {},
        "notification_action": NOTIFICATION_SKIP,
        "notification_reason": "Уведомление не меняется",
        "new_values": None,
        "ppr_event_id": event.id,
    }


def empty_summary(total_rows: int = 0) -> dict[str, Any]:
    return {
        "total_rows": total_rows,
        "valid_rows": 0,
        "invalid_rows": 0,
        "new_events": 0,
        "updated_events": 0,
        "unchanged_events": 0,
        "skipped_manual_events": 0,
        "duplicate_rows": 0,
        "missing_date_rows": 0,
        "notifications_to_create": 0,
        "notifications_to_update": 0,
        "notifications_to_skip": 0,
        "events_missing_from_excel": 0,
        "errors_count": 0,
        "matched_existing_by_source_key": 0,
        "matched_existing_by_external_id": 0,
        "matched_existing_by_legacy": 0,
        "matched_existing_by_fingerprint": 0,
        "genuinely_new": 0,
        "ambiguous": 0,
        "duplicates": 0,
        "key_by_method": {
            KEY_METHOD_ID: 0,
            KEY_METHOD_SOURCE: 0,
            KEY_METHOD_FINGERPRINT: 0,
            MATCH_METHOD_NONE: 0,
        },
        "match_by_method": {
            MATCH_METHOD_SOURCE_KEY: 0,
            MATCH_METHOD_EXTERNAL_ID: 0,
            MATCH_METHOD_LEGACY_ROW: 0,
            MATCH_METHOD_FINGERPRINT: 0,
            MATCH_METHOD_AMBIGUOUS: 0,
            MATCH_METHOD_NONE: 0,
        },
    }


def summarize_details(details: list[dict], total_rows: int) -> dict[str, Any]:
    summary = empty_summary(total_rows)
    for item in details:
        action = item["action"]
        key_method = item.get("key_method") or MATCH_METHOD_NONE
        if key_method in summary["key_by_method"]:
            summary["key_by_method"][key_method] += 1
        else:
            summary["key_by_method"][key_method] = 1
        match_method = item.get("match_method") or MATCH_METHOD_NONE
        if match_method in summary["match_by_method"]:
            summary["match_by_method"][match_method] += 1
        else:
            summary["match_by_method"][match_method] = 1
        if action not in {ACTION_CREATE, ACTION_DUPLICATE, ACTION_AMBIGUOUS, ACTION_INVALID, ACTION_MISSING_FROM_EXCEL}:
            if match_method == MATCH_METHOD_SOURCE_KEY:
                summary["matched_existing_by_source_key"] += 1
            elif match_method == MATCH_METHOD_EXTERNAL_ID:
                summary["matched_existing_by_external_id"] += 1
            elif match_method == MATCH_METHOD_LEGACY_ROW:
                summary["matched_existing_by_legacy"] += 1
            elif match_method == MATCH_METHOD_FINGERPRINT:
                summary["matched_existing_by_fingerprint"] += 1
        if action == ACTION_INVALID:
            summary["invalid_rows"] += 1
        elif action == ACTION_DUPLICATE:
            summary["duplicate_rows"] += 1
            summary["duplicates"] += 1
        elif action == ACTION_AMBIGUOUS:
            summary["ambiguous"] += 1
        elif action == ACTION_MISSING_FROM_EXCEL:
            summary["events_missing_from_excel"] += 1
            continue
        else:
            summary["valid_rows"] += 1

        if item.get("new_date") is None and action not in {ACTION_INVALID, ACTION_DUPLICATE, ACTION_MISSING_FROM_EXCEL}:
            summary["missing_date_rows"] += 1
        if action == ACTION_CREATE:
            summary["new_events"] += 1
            summary["genuinely_new"] += 1
        elif action in {ACTION_UPDATE, ACTION_SOURCE_KEY_MIGRATION}:
            summary["updated_events"] += 1
        elif action == ACTION_UNCHANGED:
            summary["unchanged_events"] += 1
        elif action == ACTION_SKIP_MANUAL:
            summary["skipped_manual_events"] += 1

        notification_action = item.get("notification_action")
        if notification_action == NOTIFICATION_CREATE:
            summary["notifications_to_create"] += 1
        elif notification_action in {NOTIFICATION_UPDATE, NOTIFICATION_DISABLE_MISSING_DATE, NOTIFICATION_DISABLE_NOTIFY_FALSE}:
            summary["notifications_to_update"] += 1
        else:
            summary["notifications_to_skip"] += 1

    summary["errors_count"] = summary["invalid_rows"] + summary["duplicate_rows"] + summary["ambiguous"]
    return summary


def build_import_preview(db: Session, content: bytes, filename: str, mode: str = IMPORT_MODE_SAFE) -> dict:
    mode = validate_mode(mode)
    file_hash = file_sha256(content)
    rows, invalid_rows, total_rows = parse_excel_content(content)
    source_key_counts: dict[str, int] = {}
    for row in rows:
        source_key_counts[row.source_key] = source_key_counts.get(row.source_key, 0) + 1
    duplicate_source_keys = {key for key, count in source_key_counts.items() if count > 1}
    seen_source_keys = {row.source_key for row in rows}
    matched_event_ids: set[int] = set()

    details: list[dict] = invalid_rows[:]
    for row in rows:
        match = find_matching_events(db, row)
        if row.source_key in duplicate_source_keys:
            if len(match.events) == 1:
                matched_event_ids.add(match.events[0].id)
            details.append(row_detail(row, match.events[0] if len(match.events) == 1 else None, ACTION_DUPLICATE, "Дублирующийся source_key/fingerprint внутри Excel", {}, mode, MATCH_METHOD_NONE, "low"))
            continue
        if len(match.events) > 1 or match.method == MATCH_METHOD_AMBIGUOUS:
            details.append(row_detail(row, None, ACTION_AMBIGUOUS, match.reason, {}, mode, MATCH_METHOD_AMBIGUOUS, "low"))
            continue

        event = match.events[0] if match.events else None
        if event is not None:
            matched_event_ids.add(event.id)
        if event is None:
            details.append(row_detail(row, None, ACTION_CREATE, "Новая ППР из Excel", {}, mode, MATCH_METHOD_NONE, "none"))
            continue

        changes = compare_event_values(event, row.values)
        only_source_key_migration = set(changes) == {"external_id", "source_key", "source_row"} or set(changes) == {"source_key"} or set(changes) == {"source_key", "source_row"}
        if mode == IMPORT_MODE_NEW_ONLY:
            details.append(row_detail(row, event, ACTION_UNCHANGED, "new_only: существующая ППР не меняется", {}, mode, match.method, "high"))
        elif event.is_manually_edited and mode == IMPORT_MODE_SAFE:
            details.append(row_detail(row, event, ACTION_SKIP_MANUAL, "ППР вручную изменена в Mini App, safe-импорт ее не перезаписывает", changes, mode, match.method, "high"))
        elif match.method in {MATCH_METHOD_EXTERNAL_ID, MATCH_METHOD_LEGACY_ROW, MATCH_METHOD_FINGERPRINT} and only_source_key_migration:
            details.append(row_detail(row, event, ACTION_SOURCE_KEY_MIGRATION, "source_key будет мигрирован на стабильный source_key", changes, mode, match.method, "medium"))
        elif changes:
            details.append(row_detail(row, event, ACTION_UPDATE, "ППР будет обновлена из Excel", changes, mode, match.method, "high" if match.method in {MATCH_METHOD_SOURCE_KEY, MATCH_METHOD_EXTERNAL_ID, MATCH_METHOD_FINGERPRINT} else "medium"))
        else:
            details.append(row_detail(row, event, ACTION_UNCHANGED, "Изменений нет", {}, mode, match.method, "high"))

    imported_events = (
        db.query(PprEvent)
        .filter(PprEvent.source_key.isnot(None))
        .filter(~PprEvent.external_id.like("MANUAL-%"))
        .all()
    )
    for event in imported_events:
        if event.id in matched_event_ids:
            continue
        key = event.source_key or source_key_from_external_id(event.external_id)
        if key and key not in seen_source_keys:
            details.append(missing_from_excel_detail(event))

    summary = summarize_details(details, total_rows)
    successful_same_hash = (
        db.query(ImportRun)
        .filter(ImportRun.file_hash == file_hash, ImportRun.status == IMPORT_STATUS_COMPLETED)
        .order_by(ImportRun.id.desc())
        .first()
    )
    warnings: list[str] = []
    if successful_same_hash:
        warnings.append("Этот файл уже был успешно импортирован. Повторное применение без force запрещено.")

    return {
        "preview_id": None,
        "filename": filename,
        "mode": mode,
        "file_hash": file_hash,
        "summary": summary,
        "details": details,
        "warnings": warnings,
    }


def save_import_preview(db: Session, content: bytes, filename: str, mode: str, user) -> dict:
    preview = build_import_preview(db, content, filename, mode)
    preview_id = uuid4().hex
    preview["preview_id"] = preview_id
    now = datetime.utcnow()
    run = ImportRun(
        preview_id=preview_id,
        filename=filename,
        mode=preview["mode"],
        started_at=now,
        completed_at=now,
        started_by_id=user.telegram_id,
        started_by_name=user_display(user),
        status=IMPORT_STATUS_PREVIEW,
        summary=preview,
        file_hash=preview["file_hash"],
    )
    db.add(run)
    db.commit()
    return preview


def value_from_json(field: str, value: Any) -> Any:
    if value in (None, ""):
        return None
    if field == "date":
        return date.fromisoformat(value)
    if field in {"start_time", "end_time"}:
        return time.fromisoformat(value)
    return value


def detail_values(detail: dict) -> dict[str, Any]:
    values = detail.get("new_values") or {}
    result = {}
    for key, value in values.items():
        result[key] = value_from_json(key, value)
    return result


def apply_values_to_event(event: PprEvent, values: dict[str, Any]) -> dict[str, tuple[Any, Any]]:
    changes: dict[str, tuple[Any, Any]] = {}
    for field, new_value in values.items():
        if not hasattr(event, field):
            continue
        old_value = getattr(event, field)
        if old_value != new_value:
            setattr(event, field, new_value)
            changes[field] = (old_value, new_value)
    if not event.is_active:
        event.ppr_status = PPR_STATUS_ARCHIVED
    elif event.ppr_status == PPR_STATUS_ARCHIVED:
        event.ppr_status = PPR_STATUS_SCHEDULED
    event.updated_at = datetime.utcnow()
    return changes


def get_event_for_detail(db: Session, detail: dict) -> PprEvent | None:
    event_id = detail.get("ppr_event_id")
    if event_id:
        event = db.get(PprEvent, event_id)
        if event:
            return event
    source_key = detail.get("source_key")
    external_id = detail.get("external_id")
    filters = []
    if source_key:
        filters.append(PprEvent.source_key == source_key)
    if external_id:
        filters.append(PprEvent.external_id == external_id)
    if not filters:
        return None
    return db.query(PprEvent).filter(or_(*filters)).one_or_none()


def add_import_audit(db: Session, event: PprEvent, action: str, user, comment: str | None = None) -> None:
    db.add(
        AuditLog(
            ppr_event_id=event.id,
            action=action,
            user_id=user.telegram_id,
            user_name=user_display(user),
            comment=comment,
        )
    )


def duplicate_successful_import_exists(db: Session, file_hash: str, exclude_id: int | None = None) -> bool:
    query = db.query(ImportRun).filter(ImportRun.file_hash == file_hash, ImportRun.status == IMPORT_STATUS_COMPLETED)
    if exclude_id is not None:
        query = query.filter(ImportRun.id != exclude_id)
    return query.first() is not None


def apply_saved_import_preview(
    db: Session,
    preview_id: str,
    content: bytes,
    mode: str,
    user,
    confirm_force: bool = False,
) -> dict:
    mode = validate_mode(mode)
    run = db.query(ImportRun).filter(ImportRun.preview_id == preview_id).one_or_none()
    if run is None or run.status != IMPORT_STATUS_PREVIEW or not run.summary:
        raise ImportPreviewNotFound("Preview not found. Run /api/import/excel/preview first.")
    if run.mode != mode:
        raise ImportPreviewError(f"Apply mode must match preview mode: {run.mode}")

    current_hash = file_sha256(content)
    if current_hash != run.file_hash:
        raise ImportFileChanged("Excel file changed after preview. Run preview again.")
    if mode == IMPORT_MODE_FORCE and not confirm_force:
        raise ImportForceConfirmationRequired("Force import requires confirm_force=true.")
    if mode != IMPORT_MODE_FORCE and duplicate_successful_import_exists(db, run.file_hash, exclude_id=run.id):
        raise ImportRepeatedFile("This file was already imported. Use force mode with explicit confirmation to reapply.")

    preview = run.summary
    if any(detail.get("action") in {ACTION_INVALID, ACTION_DUPLICATE, ACTION_AMBIGUOUS} for detail in preview.get("details", [])):
        raise ImportPreviewError("Preview contains invalid, duplicate, or ambiguous rows. Fix the Excel file and run preview again.")

    applied = {
        "created": 0,
        "updated": 0,
        "unchanged": 0,
        "skipped": 0,
        "notifications_synced": 0,
    }
    now = datetime.utcnow()
    run.started_at = now
    run.status = "applying"
    run.error_message = None

    try:
        for detail in preview.get("details", []):
            action = detail.get("action")
            if action in {ACTION_INVALID, ACTION_DUPLICATE, ACTION_AMBIGUOUS, ACTION_SKIP_MANUAL, ACTION_MISSING_FROM_EXCEL}:
                applied["skipped"] += 1
                continue
            values = detail_values(detail)
            if not values:
                applied["skipped"] += 1
                continue

            event = get_event_for_detail(db, detail)
            if action == ACTION_CREATE:
                event = PprEvent(
                    external_id=values["external_id"],
                    source_key=values["source_key"],
                    title=values["title"],
                    ppr_status=PPR_STATUS_SCHEDULED,
                    is_manually_edited=False,
                    manual_updated_at=None,
                )
                db.add(event)
                db.flush()
                apply_values_to_event(event, values)
                add_import_audit(db, event, "import_created", user, f"Excel import {preview_id}")
                applied["created"] += 1
            elif event is None:
                applied["skipped"] += 1
                continue
            else:
                changes = apply_values_to_event(event, values) if action in {ACTION_UPDATE, ACTION_SOURCE_KEY_MIGRATION} else {}
                if changes:
                    event.is_manually_edited = False
                    event.manual_updated_at = None
                    add_import_audit(db, event, "import_updated", user, describe_changes(changes))
                    applied["updated"] += 1
                else:
                    applied["unchanged"] += 1

            if event is not None and detail.get("notification_action") != NOTIFICATION_SKIP:
                sync_start_notification(db, event, user, audit=True)
                applied["notifications_synced"] += 1

        completed_summary = {
            **preview,
            "applied": applied,
        }
        run.status = IMPORT_STATUS_COMPLETED
        run.completed_at = datetime.utcnow()
        run.started_by_id = user.telegram_id
        run.started_by_name = user_display(user)
        run.summary = completed_summary
        db.commit()
        return completed_summary
    except Exception as exc:
        db.rollback()
        failed_run = db.query(ImportRun).filter(ImportRun.preview_id == preview_id).one_or_none()
        if failed_run:
            failed_run.status = IMPORT_STATUS_FAILED
            failed_run.completed_at = datetime.utcnow()
            failed_run.error_message = str(exc)
            db.commit()
        raise
