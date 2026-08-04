from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from app.config import get_settings
from app.db.models import PprEvent
from app.excel.importer import HEADER_ALIASES
from app.services.statuses import PPR_STATUS_ARCHIVED
from app.services.user_service import user_display


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MAIN_SHEET_NAME = "ППР"
INFO_SHEET_NAME = "Информация"
SERVICE_SHEET_NAME = "Служебная информация"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOP_ALIGNMENT = Alignment(vertical="top")
WRAPPED_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

MAIN_COLUMN_WIDTHS = {
    "ID": 22,
    "Дата": 13,
    "Время выхода": 16,
    "Время завершения": 18,
    "Тип уведомления": 19,
    "Проект": 24,
    "Название ППР": 42,
    "Активности": 55,
    "Ответственный настройка": 28,
    "Ответственный отчетка": 28,
    "Ссылка": 38,
    "Уведомлять о выходе": 22,
    "Уведомлять о завершении": 25,
    "Активно": 12,
    "Комментарий": 45,
    "Исходная строка": 18,
}

WRAPPED_FIELDS = {
    "title",
    "activities",
    "responsible_setup",
    "responsible_report",
    "source_link",
    "comment",
}


@dataclass(frozen=True)
class PprExcelExport:
    content: bytes
    filename: str
    event_count: int
    generated_at: datetime


def _set_text_cell(cell, value: str) -> None:
    """Keep user text as text even when it starts with an Excel formula marker."""
    cell.value = value
    cell.data_type = "s"


def _style_header(worksheet) -> None:
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
    worksheet.row_dimensions[1].height = 32
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(worksheet.max_column)}{max(worksheet.max_row, 1)}"


def _append_main_sheet(workbook: Workbook, events: list[PprEvent]) -> None:
    worksheet = workbook.active
    worksheet.title = MAIN_SHEET_NAME
    headers = list(HEADER_ALIASES)
    worksheet.append(headers)

    for event in events:
        row_index = worksheet.max_row + 1
        for column_index, header in enumerate(headers, start=1):
            field = HEADER_ALIASES[header]
            value = getattr(event, field)
            cell = worksheet.cell(row=row_index, column=column_index)
            if isinstance(value, str):
                _set_text_cell(cell, value)
            else:
                cell.value = value

    for row in worksheet.iter_rows(min_row=2):
        for header, cell in zip(headers, row):
            field = HEADER_ALIASES[header]
            cell.alignment = WRAPPED_ALIGNMENT if field in WRAPPED_FIELDS else TOP_ALIGNMENT
            if field == "external_id":
                cell.number_format = "@"
            elif field == "date":
                cell.number_format = "DD.MM.YYYY"
            elif field in {"start_time", "end_time"}:
                cell.number_format = "HH:MM:SS"

    for column_index, header in enumerate(headers, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = MAIN_COLUMN_WIDTHS[header]
    _style_header(worksheet)


def _event_source(event: PprEvent) -> str:
    if event.external_id.startswith("MANUAL-"):
        return "Ручное создание"
    return "Excel-импорт"


def _append_service_sheet(workbook: Workbook, events: list[PprEvent]) -> None:
    worksheet = workbook.create_sheet(SERVICE_SHEET_NAME)
    headers = [
        "ID ППР",
        "Текущий статус ППР",
        "Количество уведомлений",
        "Статусы уведомлений",
        "Источник создания",
        "Изменена вручную",
    ]
    worksheet.append(headers)
    for event in events:
        notifications = sorted(event.notifications, key=lambda item: (item.type, item.id))
        statuses = "; ".join(f"{item.type}: {item.status}" for item in notifications)
        worksheet.append(
            [
                event.external_id,
                PPR_STATUS_ARCHIVED if not event.is_active else event.ppr_status,
                len(notifications),
                statuses,
                _event_source(event),
                event.is_manually_edited,
            ]
        )

    for row in worksheet.iter_rows(min_row=2):
        _set_text_cell(row[0], str(row[0].value))
        for cell in row:
            cell.alignment = WRAPPED_ALIGNMENT
    for index, width in enumerate((22, 24, 24, 45, 22, 22), start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    _style_header(worksheet)


def _append_info_sheet(workbook: Workbook, events: list[PprEvent], user, generated_at: datetime) -> None:
    worksheet = workbook.create_sheet(INFO_SHEET_NAME)
    worksheet.append(["Параметр", "Значение"])
    archived_count = sum(1 for event in events if not event.is_active or event.ppr_status == PPR_STATUS_ARCHIVED)
    rows = [
        ("Дата и время формирования", generated_at.strftime("%d.%m.%Y %H:%M:%S %Z")),
        ("Администратор", user_display(user)),
        ("Telegram ID", user.telegram_id),
        ("Количество карточек", len(events)),
        ("Карточек без даты", sum(1 for event in events if event.date is None)),
        ("Архивных карточек", archived_count),
        ("Назначение", "Массовое редактирование текущей базы ППР с безопасной обратной загрузкой."),
        (
            "Инструкция",
            "После редактирования загрузите файл во вкладке Импорт, выполните Preview, проверьте изменения и только после этого нажмите Apply",
        ),
        ("Предупреждение", "Не удаляйте и не изменяйте столбец ID"),
        ("Предупреждение", "Не изменяйте названия столбцов"),
        (
            "Предупреждение",
            "Статусы взятия в работу, проверки и отправки уведомлений данным файлом не сбрасываются",
        ),
    ]
    for label, value in rows:
        worksheet.append([label, value])

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAPPED_ALIGNMENT
            if isinstance(cell.value, str):
                cell.data_type = "s"
    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 100
    _style_header(worksheet)


def build_ppr_excel_export(db: Session, user, now: datetime | None = None) -> PprExcelExport:
    timezone = ZoneInfo(get_settings().default_timezone)
    if now is None:
        generated_at = datetime.now(timezone)
    elif now.tzinfo is None:
        generated_at = now.replace(tzinfo=timezone)
    else:
        generated_at = now.astimezone(timezone)

    events = (
        db.query(PprEvent)
        .options(selectinload(PprEvent.notifications))
        .order_by(PprEvent.id.asc())
        .all()
    )
    workbook = Workbook()
    _append_main_sheet(workbook, events)
    _append_service_sheet(workbook, events)
    _append_info_sheet(workbook, events, user, generated_at)
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False

    output = BytesIO()
    workbook.save(output)
    filename = f"PPR_export_{generated_at.strftime('%Y-%m-%d_%H-%M')}.xlsx"
    return PprExcelExport(
        content=output.getvalue(),
        filename=filename,
        event_count=len(events),
        generated_at=generated_at,
    )
