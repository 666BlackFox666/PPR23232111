"""Safe, all-or-nothing replacement of the Excel-managed schedule.

This is intentionally separate from the incremental Excel import service.  It
uses that service's parser and source-key rules, but has stricter validation:
the input must describe a complete future schedule before any existing PPR
data can be removed.
"""

from __future__ import annotations

import shutil
import subprocess
from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, text
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models import AppUser, AuditLog, ImportRun, PprEvent, PprNotification
from app.excel.import_service import (
    KEY_METHOD_FINGERPRINT,
    ParsedExcelRow,
    automatic_external_id,
    build_fingerprint_source_key,
    file_sha256,
    parse_excel_content,
)
from app.excel.importer import cell, get_header_map
from app.services.statuses import (
    NOTIFICATION_STATUS_PLANNED,
    NOTIFICATION_STATUS_PROCESSING,
    NOTIFICATION_STATUSES,
    PPR_STATUS_IN_PROGRESS,
    PPR_STATUS_SCHEDULED,
    PPR_STATUS_VERIFIED,
)
from app.services.telegram_sender import scheduler_status


REPLACE_CONFIRMATION = "REPLACE_SCHEDULE"
ADVISORY_LOCK_KEY = 8_830_914_217
REQUIRED_SCHEDULE_COLUMNS = ("title", "date", "start_time")
TRUE_FALSE_VALUES = {"да", "yes", "y", "true", "1", "истина", "нет", "no", "n", "false", "0", "ложь"}


class ScheduleReplacementError(RuntimeError):
    pass


class ScheduleValidationError(ScheduleReplacementError):
    def __init__(self, preview: dict[str, Any]):
        super().__init__("Excel schedule validation failed")
        self.preview = preview


def project_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _moscow_now(now: datetime | None = None) -> datetime:
    if now is not None:
        return now.replace(tzinfo=None)
    return datetime.now(ZoneInfo("Europe/Moscow")).replace(tzinfo=None)


def _is_bool_value(value: Any) -> bool:
    if value in (None, "") or isinstance(value, bool):
        return True
    return str(value).strip().lower() in TRUE_FALSE_VALUES


def _append_error(errors: list[dict[str, Any]], row: int | None, reason: str) -> None:
    entry = {"excel_row_number": row, "reason": reason}
    if entry not in errors:
        errors.append(entry)


def _current_db_summary(db: Session) -> dict[str, Any]:
    notification_statuses = dict(db.query(PprNotification.status, func.count(PprNotification.id)).group_by(PprNotification.status).all())
    ppr_statuses = dict(db.query(PprEvent.ppr_status, func.count(PprEvent.id)).group_by(PprEvent.ppr_status).all())
    return {
        "ppr_events": db.query(PprEvent).count(),
        "ppr_notifications": db.query(PprNotification).count(),
        "planned": int(notification_statuses.get(NOTIFICATION_STATUS_PLANNED, 0)),
        "sent": int(notification_statuses.get("sent", 0)),
        "verified": int(ppr_statuses.get(PPR_STATUS_VERIFIED, 0)),
        "in_progress": int(ppr_statuses.get(PPR_STATUS_IN_PROGRESS, 0)),
        "notifications_by_status": {str(key): int(value) for key, value in notification_statuses.items()},
        "ppr_by_status": {str(key): int(value) for key, value in ppr_statuses.items()},
    }


def _build_replacement_rows(rows: list[ParsedExcelRow], sheet_name: str) -> list[ParsedExcelRow]:
    """Give each row a stable, unique identity for full replacement only.

    Incremental import intentionally keeps its fingerprint matching logic.  A
    replacement has no existing rows to match, so a no-ID row must keep its
    own Excel-row identity rather than collide with an equal fingerprint.
    """
    for row in rows:
        if row.external_id_from_excel:
            source_key = f"id:{row.external_id_from_excel}"
        elif row.source_value:
            source_key = f"source:{row.source_value}"
        else:
            source_key = f"replace-row:{sheet_name}:{row.excel_row_number}"
        row.source_key = source_key
        row.values["source_key"] = source_key
        if not row.external_id_from_excel:
            row.external_id = automatic_external_id(source_key)
            row.values["external_id"] = row.external_id
    return rows


def build_schedule_replacement_preview(
    db: Session,
    content: bytes,
    filename: str,
    *,
    now: datetime | None = None,
) -> dict[str, Any]:
    """Validate a complete schedule without changing the database."""
    errors: list[dict[str, Any]] = []
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
        worksheet = workbook["ППР_для_бота"] if "ППР_для_бота" in workbook.sheetnames else workbook.active
        header_map = get_header_map(worksheet)
    except Exception as exc:
        return {
            "filename": filename,
            "sha256": file_sha256(content),
            "current_db": _current_db_summary(db),
            "new_excel": {"total_rows": 0, "unique_ppr": 0, "future_ppr": 0, "notifications_to_create": 0, "min_date": None, "max_date": None, "duplicates": 0, "errors": 1},
            "errors": [{"excel_row_number": None, "reason": f"Не удалось открыть Excel: {exc}"}],
            "result": {},
        }

    for field in REQUIRED_SCHEDULE_COLUMNS:
        if field not in header_map:
            _append_error(errors, None, f"В Excel нет обязательной колонки '{field}'")

    rows, parser_errors, total_rows = parse_excel_content(content)
    rows = _build_replacement_rows(rows, worksheet.title)
    for item in parser_errors:
        _append_error(errors, item.get("excel_row_number"), item.get("reason", "Строка невалидна"))
    if total_rows == 0:
        _append_error(errors, None, "Excel не содержит ни одной ППР")

    excel_now = _moscow_now(now)
    ids: dict[str, list[int]] = defaultdict(list)
    source_keys: dict[str, list[int]] = defaultdict(list)
    fingerprints: dict[str, list[int]] = defaultdict(list)
    notification_count = 0
    future_rows = 0
    dates = []

    for row in rows:
        values = row.values
        raw_date = cell(tuple(worksheet[row.excel_row_number]), header_map, "date") if "date" in header_map else None
        raw_time = cell(tuple(worksheet[row.excel_row_number]), header_map, "start_time") if "start_time" in header_map else None
        # A card without a date is valid and intentionally has no time or
        # notification.  When a date exists, however, sending time remains a
        # required part of the scheduled PPR contract.
        if values.get("date") is not None and raw_time in (None, ""):
            _append_error(errors, row.excel_row_number, "Время выхода не заполнено")
        if values.get("date") is not None and values.get("start_time") is None:
            _append_error(errors, row.excel_row_number, "Время выхода не распознано")

        for field in ("notify_start", "notify_end", "is_active"):
            raw_value = cell(tuple(worksheet[row.excel_row_number]), header_map, field) if field in header_map else None
            if not _is_bool_value(raw_value):
                _append_error(errors, row.excel_row_number, f"Некорректный флаг '{field}': {raw_value}")

        if row.external_id_from_excel:
            ids[row.external_id_from_excel].append(row.excel_row_number)
        source_keys[row.source_key].append(row.excel_row_number)
        if row.key_method == KEY_METHOD_FINGERPRINT:
            fingerprints[build_fingerprint_source_key(values)].append(row.excel_row_number)

        if values.get("date"):
            dates.append(values["date"])
            if values.get("start_time") and raw_time not in (None, ""):
                scheduled_at = datetime.combine(values["date"], values["start_time"])
                if scheduled_at <= excel_now:
                    _append_error(errors, row.excel_row_number, f"Дата и время ППР должны быть в будущем ({scheduled_at.isoformat(sep=' ')})")
                else:
                    future_rows += 1
                    if values.get("is_active", True) and values.get("notify_start", True):
                        notification_count += 1

    duplicate_groups = 0
    warnings: list[dict[str, Any]] = []
    for external_id, row_numbers in ids.items():
        if len(row_numbers) > 1:
            duplicate_groups += 1
            for row_number in row_numbers:
                _append_error(errors, row_number, f"Дублирующийся Excel ID: {external_id}")
    for source_key, row_numbers in source_keys.items():
        if len(row_numbers) > 1:
            duplicate_groups += 1
            for row_number in row_numbers:
                _append_error(errors, row_number, f"Дублирующийся source_key: {source_key}")
    for fingerprint, row_numbers in fingerprints.items():
        if len(row_numbers) > 1:
            warnings.append({
                "type": "duplicate_fingerprint",
                "fingerprint": fingerprint,
                "excel_row_numbers": sorted(row_numbers),
                "reason": "Совпадающий fingerprint допустим в режиме полной замены: будут созданы отдельные ППР.",
            })

    current = _current_db_summary(db)
    new_excel = {
        "total_rows": total_rows,
        "unique_ppr": len(rows),
        "future_ppr": future_rows,
        "dated_ppr": sum(1 for row in rows if row.values.get("date") is not None),
        "missing_date_ppr": sum(1 for row in rows if row.values.get("date") is None),
        "notifications_to_create": notification_count,
        "min_date": min(dates).isoformat() if dates else None,
        "max_date": max(dates).isoformat() if dates else None,
        "duplicate_excel_id_groups": sum(1 for values in ids.values() if len(values) > 1),
        "duplicate_source_key_groups": sum(1 for values in source_keys.values() if len(values) > 1),
        "ambiguous_fingerprint_groups": sum(1 for values in fingerprints.values() if len(values) > 1),
        "duplicates": duplicate_groups,
        "errors": len(errors),
        "warnings": len(warnings),
    }
    return {
        "filename": filename,
        "sha256": file_sha256(content),
        "current_db": {**current, "ppr_events_to_delete": current["ppr_events"], "notifications_to_delete": current["ppr_notifications"]},
        "new_excel": new_excel,
        "result": {
            "ppr_events_to_delete": current["ppr_events"],
            "notifications_to_delete": current["ppr_notifications"],
            "ppr_events_to_create": len(rows),
            "notifications_to_create": notification_count,
            "users_preserved": True,
            "roles_preserved": True,
            "audit_log_preserved": True,
            "sequences_reset": False,
        },
        "errors": errors,
        "warnings": warnings,
    }


def _diagnostic_value(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return value


def _diagnostic_row_data(row: ParsedExcelRow | None, worksheet, header_map: dict[str, int], row_number: int) -> dict[str, Any]:
    values = row.values if row else {}
    def raw(field: str) -> Any:
        return cell(tuple(worksheet[row_number]), header_map, field) if field in header_map else None
    return {
        "excel_row_number": row_number,
        # Only the literal Excel ID belongs in the diagnostics.  AUTO-* is an
        # importer-generated external_id and must not be reported as an Excel
        # ID duplicate.
        "external_id": raw("external_id"),
        "project": values.get("project"),
        "title": values.get("title") or raw("title"),
        "date": _diagnostic_value(values.get("date") or raw("date")),
        "start_time": _diagnostic_value(values.get("start_time") or raw("start_time")),
        "end_time": _diagnostic_value(values.get("end_time") or raw("end_time")),
        "source_key": row.source_key if row else None,
        "fingerprint": build_fingerprint_source_key(values) if row else None,
        "key_method": row.key_method if row else None,
        "values": values,
    }


def _recommended_action(reasons: list[str]) -> str:
    text_reasons = " ".join(reasons).lower()
    actions = []
    if "дата" in text_reasons:
        actions.append("заполнить дату")
    if "время" in text_reasons:
        actions.append("заполнить время")
    if "excel id" in text_reasons:
        actions.append("изменить ID")
    if "source_key" in text_reasons or "fingerprint" in text_reasons:
        actions.append("проверить дубликат")
    if not actions:
        actions.append("удалить строку")
    return "; ".join(dict.fromkeys(actions))


def export_validation_report(content: bytes, output_path: str | Path, preview: dict[str, Any]) -> Path:
    """Write a diagnostic workbook without opening a database write transaction."""
    workbook = load_workbook(BytesIO(content), data_only=True)
    worksheet = workbook["ППР_для_бота"] if "ППР_для_бота" in workbook.sheetnames else workbook.active
    header_map = get_header_map(worksheet)
    rows, parser_errors, _ = parse_excel_content(content)
    rows = _build_replacement_rows(rows, worksheet.title)
    rows_by_number = {row.excel_row_number: row for row in rows}
    errors_by_row: dict[int, list[str]] = defaultdict(list)
    for item in preview.get("errors", []):
        row_number = item.get("excel_row_number")
        if row_number is not None:
            errors_by_row[int(row_number)].append(str(item.get("reason", "Ошибка")))
    for item in parser_errors:
        row_number = item.get("excel_row_number")
        if row_number is not None:
            errors_by_row[int(row_number)].append(str(item.get("reason", "Ошибка")))

    diagnostic_rows = {
        row_number: _diagnostic_row_data(rows_by_number.get(row_number), worksheet, header_map, row_number)
        for row_number in set(errors_by_row) | set(rows_by_number)
    }

    error_headers = [
        "Номер строки Excel", "Excel ID", "Проект", "Название ППР", "Дата", "Время выхода",
        "Время завершения", "source_key", "fingerprint", "Причина ошибки", "Рекомендуемое действие",
    ]
    duplicate_headers = [
        "Номер группы", "Тип конфликта", "source_key/fingerprint", "Номера строк", "Excel ID",
        "Проект", "Название ППР", "Дата и время", "Какие поля отличаются", "Какие поля полностью совпадают",
    ]
    valid_headers = [
        "Номер строки Excel", "Excel ID", "Проект", "Название ППР", "Дата", "Время выхода",
        "Время завершения", "source_key", "fingerprint", "Уведомлять о выходе", "Активно",
    ]
    missing_date_headers = ["Строка Excel", "Excel ID", "Проект", "Название ППР", "Активности", "source_key", "Причина"]

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    report = Workbook()
    summary_sheet = report.active
    summary_sheet.title = "Summary"
    errors_sheet = report.create_sheet("Errors")
    duplicates_sheet = report.create_sheet("DuplicateGroups")
    missing_date_sheet = report.create_sheet("Без даты")
    valid_sheet = report.create_sheet("ValidRows")

    summary = preview.get("new_excel", {})
    summary_rows = [
        ("Всего ППР", summary.get("total_rows", 0)),
        ("ППР с датой", summary.get("dated_ppr", 0)),
        ("ППР без даты", summary.get("missing_date_ppr", 0)),
        ("Уведомлений, которые будут созданы", summary.get("notifications_to_create", 0)),
        ("Повторяющихся fingerprint", summary.get("ambiguous_fingerprint_groups", 0)),
        ("Блокирующих ошибок", summary.get("errors", 0)),
        ("Предупреждений", summary.get("warnings", 0)),
        ("Дублей Excel ID", summary.get("duplicate_excel_id_groups", 0)),
        ("Дублей source_key", summary.get("duplicate_source_key_groups", 0)),
        ("Минимальная дата", summary.get("min_date")),
        ("Максимальная дата", summary.get("max_date")),
        ("SHA-256 файла", preview.get("sha256")),
    ]
    summary_sheet.append(["Показатель", "Значение"])
    for item in summary_rows:
        summary_sheet.append(list(item))

    errors_sheet.append(error_headers)
    for row_number in sorted(errors_by_row):
        item = diagnostic_rows[row_number]
        reasons = list(dict.fromkeys(errors_by_row[row_number]))
        errors_sheet.append([
            item["excel_row_number"], item["external_id"], item["project"], item["title"], item["date"],
            item["start_time"], item["end_time"], item["source_key"], item["fingerprint"],
            "; ".join(reasons), _recommended_action(reasons),
        ])

    duplicate_groups: dict[tuple[int, ...], dict[str, Any]] = {}
    rows_by_source = defaultdict(list)
    rows_by_fingerprint = defaultdict(list)
    rows_by_id = defaultdict(list)
    for item in diagnostic_rows.values():
        if item["source_key"]:
            rows_by_source[item["source_key"]].append(item)
        if item["fingerprint"] and item.get("key_method") == KEY_METHOD_FINGERPRINT:
            rows_by_fingerprint[item["fingerprint"]].append(item)
        if item["external_id"] and str(item["external_id"]).strip():
            rows_by_id[str(item["external_id"])].append(item)

    def add_group(conflict_type: str, key: str, group_rows: list[dict[str, Any]]) -> None:
        if len(group_rows) < 2:
            return
        numbers = tuple(sorted(item["excel_row_number"] for item in group_rows))
        group = duplicate_groups.setdefault(numbers, {"types": [], "keys": [], "rows": group_rows})
        if conflict_type not in group["types"]:
            group["types"].append(conflict_type)
        if key not in group["keys"]:
            group["keys"].append(key)

    for key, group_rows in rows_by_id.items():
        add_group("duplicate Excel ID", key, group_rows)
    for key, group_rows in rows_by_source.items():
        add_group("duplicate source_key", key, group_rows)
    for key, group_rows in rows_by_fingerprint.items():
        add_group("повторяющийся fingerprint (предупреждение)", key, group_rows)

    duplicate_groups_sheet = duplicates_sheet
    duplicate_groups_sheet.append(duplicate_headers)
    for group_number, group in enumerate(duplicate_groups.values(), start=1):
        group_rows = group["rows"]
        fields = ["title", "project", "date", "start_time", "end_time", "activities", "notification_type", "notify_start", "is_active"]
        different, same = [], []
        for field in fields:
            values = {_diagnostic_value(item["values"].get(field)) for item in group_rows}
            (same if len(values) == 1 else different).append(field)
        duplicate_groups_sheet.append([
            group_number, ", ".join(group["types"]), ", ".join(group["keys"]),
            ", ".join(str(item["excel_row_number"]) for item in group_rows),
            ", ".join(str(item["external_id"] or "") for item in group_rows),
            ", ".join(str(item["project"] or "") for item in group_rows),
            ", ".join(str(item["title"] or "") for item in group_rows),
            ", ".join(f"{item['date']} {item['start_time']}" for item in group_rows),
            ", ".join(different) or "нет", ", ".join(same) or "нет",
        ])

    missing_date_sheet.append(missing_date_headers)
    for row_number in sorted(rows_by_number):
        item = diagnostic_rows[row_number]
        if item["values"].get("date") is None:
            missing_date_sheet.append([
                item["excel_row_number"], item["external_id"], item["project"], item["title"],
                item["values"].get("activities"), item["source_key"], "Дата ещё не назначена",
            ])

    valid_sheet.append(valid_headers)
    for row_number in sorted(rows_by_number):
        if row_number in errors_by_row:
            continue
        item = diagnostic_rows[row_number]
        values = item["values"]
        valid_sheet.append([
            item["excel_row_number"], item["external_id"], item["project"], item["title"], item["date"],
            item["start_time"], item["end_time"], item["source_key"], item["fingerprint"],
            values.get("notify_start"), values.get("is_active"),
        ])

    error_fill = PatternFill("solid", fgColor="F4CCCC")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in report.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell_obj in sheet[1]:
            cell_obj.font = Font(bold=True)
            cell_obj.fill = header_fill
        for column_cells in sheet.columns:
            width = min(max(len(str(cell_obj.value or "")) for cell_obj in column_cells) + 2, 48)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    for row_number in range(2, errors_sheet.max_row + 1):
        for cell_obj in errors_sheet[row_number]:
            cell_obj.fill = error_fill
    report.save(output)
    return output


def _try_transaction_lock(db: Session) -> bool:
    if db.bind is None or db.bind.dialect.name != "postgresql":
        return True
    return bool(db.execute(text("SELECT pg_try_advisory_xact_lock(:lock_key)"), {"lock_key": ADVISORY_LOCK_KEY}).scalar())


def assert_replacement_runtime_is_safe(db: Session) -> None:
    settings = get_settings()
    if settings.notifications_auto_send_enabled:
        raise ScheduleReplacementError("NOTIFICATIONS_AUTO_SEND_ENABLED must be false before replacing schedule")
    processing = db.query(PprNotification).filter(PprNotification.status == NOTIFICATION_STATUS_PROCESSING).count()
    if processing:
        raise ScheduleReplacementError(f"Replacement blocked: {processing} notification(s) are processing")
    # A bot-only stop can leave a fresh heartbeat in PostgreSQL until its TTL
    # expires.  With global auto-send disabled it cannot deliver anything; the
    # processing check above still prevents replacement during an in-flight
    # delivery.
    status = scheduler_status(db)
    if status["scheduler_running"] and settings.notifications_auto_send_enabled:
        raise ScheduleReplacementError("Replacement blocked: scheduler is running")


def _docker_backup(project: Path) -> Path:
    docker = shutil.which("docker")
    if not docker:
        raise ScheduleReplacementError("Database backup failed: docker was not found")
    health = subprocess.run([docker, "inspect", "-f", "{{.State.Health.Status}}", "pprbot-postgres"], cwd=project, capture_output=True, text=True, check=False)
    if health.returncode != 0 or health.stdout.strip() != "healthy":
        raise ScheduleReplacementError("Database backup failed: PostgreSQL container is not healthy")
    backup_dir = project / "backups"
    backup_dir.mkdir(exist_ok=True)
    filename = f"pprbot-{datetime.now().strftime('%Y%m%d-%H%M%S')}.dump"
    destination = backup_dir / filename
    container_path = f"/tmp/{filename}"
    try:
        dumped = subprocess.run([docker, "exec", "pprbot-postgres", "pg_dump", "-U", "ppr_user", "-d", "ppr_db", "-Fc", "-f", container_path], cwd=project, capture_output=True, text=True, check=False)
        if dumped.returncode != 0:
            raise ScheduleReplacementError("Database backup failed: pg_dump returned an error")
        copied = subprocess.run([docker, "cp", f"pprbot-postgres:{container_path}", str(destination)], cwd=project, capture_output=True, text=True, check=False)
        if copied.returncode != 0 or not destination.exists() or destination.stat().st_size == 0:
            raise ScheduleReplacementError("Database backup failed: backup file was not created")
        return destination
    finally:
        subprocess.run([docker, "exec", "pprbot-postgres", "rm", "-f", container_path], cwd=project, capture_output=True, text=True, check=False)


def create_schedule_backup(project: Path | None = None) -> Path:
    """Use the same Docker pg_dump format as scripts/backup-db.ps1."""
    return _docker_backup(project or project_root())


def _create_events_from_rows(db: Session, rows: list[ParsedExcelRow]) -> tuple[int, int]:
    created_events = 0
    created_notifications = 0
    for row in rows:
        values = row.values
        event = PprEvent(
            external_id=values["external_id"],
            source_key=values["source_key"],
            source_row=values["source_row"],
            date=values["date"],
            start_time=values["start_time"],
            end_time=values["end_time"],
            notification_type=values["notification_type"],
            project=values["project"],
            title=values["title"],
            activities=values["activities"],
            responsible_setup=values["responsible_setup"],
            responsible_report=values["responsible_report"],
            source_link=values["source_link"],
            notify_start=values["notify_start"],
            notify_end=values["notify_end"],
            is_active=values["is_active"],
            ppr_status=PPR_STATUS_SCHEDULED,
            is_manually_edited=False,
            manual_updated_at=None,
            comment=values["comment"],
        )
        db.add(event)
        db.flush()
        created_events += 1
        if event.date and event.start_time and event.is_active and event.notify_start:
            db.add(PprNotification(
                ppr_event_id=event.id,
                type="start",
                scheduled_at=datetime.combine(event.date, event.start_time),
                status=NOTIFICATION_STATUS_PLANNED,
                auto_send_enabled=True,
                attempt_count=0,
            ))
            created_notifications += 1
    return created_events, created_notifications


def _detach_schedule_audit(db: Session) -> None:
    event_ids = [event_id for (event_id,) in db.query(PprEvent.id).all()]
    notification_ids = [notification_id for (notification_id,) in db.query(PprNotification.id).all()]
    if notification_ids:
        db.query(AuditLog).filter(AuditLog.notification_id.in_(notification_ids)).update({AuditLog.notification_id: None}, synchronize_session=False)
    if event_ids:
        db.query(AuditLog).filter(AuditLog.ppr_event_id.in_(event_ids)).update({AuditLog.ppr_event_id: None}, synchronize_session=False)


def _verify_replacement(
    db: Session,
    expected_events: int,
    expected_notifications: int,
    expected_missing_date: int,
    old_event_max_id: int | None,
    old_notification_max_id: int | None,
) -> None:
    if db.query(PprEvent).count() != expected_events or db.query(PprNotification).count() != expected_notifications:
        raise ScheduleReplacementError("Post-apply count verification failed")
    if db.query(PprEvent).filter(PprEvent.date.is_(None)).count() != expected_missing_date:
        raise ScheduleReplacementError("Post-apply validation failed: unexpected missing-date count")
    if db.query(PprEvent.source_key, func.count(PprEvent.id)).group_by(PprEvent.source_key).having(func.count(PprEvent.id) > 1).count():
        raise ScheduleReplacementError("Post-apply validation failed: duplicate source_key")
    invalid_notifications = db.query(PprNotification).filter(~PprNotification.status.in_(NOTIFICATION_STATUSES)).count()
    if invalid_notifications or db.query(PprNotification).filter(PprNotification.status != NOTIFICATION_STATUS_PLANNED).count():
        raise ScheduleReplacementError("Post-apply validation failed: invalid notification status")
    if db.query(PprNotification).filter(PprNotification.scheduled_at.is_(None)).count():
        raise ScheduleReplacementError("Post-apply validation failed: notification without scheduled_at")
    if db.query(PprNotification).filter(PprNotification.status == NOTIFICATION_STATUS_PROCESSING).count():
        raise ScheduleReplacementError("Post-apply validation failed: processing notification")
    now = _moscow_now()
    if db.query(PprNotification).filter(PprNotification.scheduled_at <= now).count():
        raise ScheduleReplacementError("Post-apply validation failed: notification in the past")
    if db.bind and db.bind.dialect.name == "postgresql":
        if old_event_max_id is not None and db.query(func.min(PprEvent.id)).scalar() <= old_event_max_id:
            raise ScheduleReplacementError("Post-apply validation failed: PPR sequence was reused")
        if old_notification_max_id is not None and db.query(func.min(PprNotification.id)).scalar() <= old_notification_max_id:
            raise ScheduleReplacementError("Post-apply validation failed: notification sequence was reused")


def replace_schedule_from_excel(
    db: Session,
    content: bytes,
    filename: str,
    *,
    confirmation: str,
    backup_creator: Callable[[], Path] | None = None,
) -> dict[str, Any]:
    if confirmation != REPLACE_CONFIRMATION:
        raise ScheduleReplacementError(f"Apply requires --confirm {REPLACE_CONFIRMATION}")

    preview = build_schedule_replacement_preview(db, content, filename)
    if preview["errors"]:
        raise ScheduleValidationError(preview)

    users_before = db.query(AppUser).count()
    roles_before = Counter(role for (role,) in db.query(AppUser.role).all())
    backup_creator = backup_creator or create_schedule_backup
    # SQLAlchemy opens a transaction for the preview queries above.  End that
    # read-only transaction before starting the transaction that owns the
    # PostgreSQL advisory lock and the replacement itself.
    db.rollback()
    try:
        with db.begin():
            if not _try_transaction_lock(db):
                raise ScheduleReplacementError("Replacement already running (PostgreSQL advisory lock is busy)")
            assert_replacement_runtime_is_safe(db)
            # Revalidate while holding the lock; the file and the database may have changed after preview.
            preview = build_schedule_replacement_preview(db, content, filename)
            if preview["errors"]:
                raise ScheduleValidationError(preview)
            backup_path = backup_creator()
            if not backup_path or not Path(backup_path).exists() or Path(backup_path).stat().st_size == 0:
                raise ScheduleReplacementError("Database backup failed: backup file was not created")

            old_event_count = db.query(PprEvent).count()
            old_notification_count = db.query(PprNotification).count()
            old_event_max_id = db.query(func.max(PprEvent.id)).scalar()
            old_notification_max_id = db.query(func.max(PprNotification.id)).scalar()
            _detach_schedule_audit(db)
            db.query(PprNotification).delete(synchronize_session=False)
            db.query(PprEvent).delete(synchronize_session=False)
            db.query(ImportRun).delete(synchronize_session=False)
            # Bulk deletes intentionally bypass ORM cascades.  Remove deleted
            # instances from this session before PostgreSQL allocates new IDs;
            # this also keeps SQLite test sessions from retaining stale rows.
            db.expunge_all()

            rows, parser_errors, _ = parse_excel_content(content)
            rows = _build_replacement_rows(rows, "ППР_для_бота")
            if parser_errors:
                raise ScheduleReplacementError("Excel changed or became invalid during replacement")
            created_events, created_notifications = _create_events_from_rows(db, rows)
            comment = (
                f"filename={filename}; sha256={preview['sha256']}; deleted_ppr={old_event_count}; "
                f"deleted_notifications={old_notification_count}; created_ppr={created_events}; "
                f"created_notifications={created_notifications}; backup={Path(backup_path)}"
            )
            db.add(AuditLog(action="schedule_replaced_from_excel", user_id="system", user_name="CLI", comment=comment))
            db.flush()
            _verify_replacement(
                db,
                created_events,
                created_notifications,
                preview["new_excel"]["missing_date_ppr"],
                old_event_max_id,
                old_notification_max_id,
            )
    except Exception:
        db.rollback()
        raise

    users_after = db.query(AppUser).count()
    roles_after = Counter(role for (role,) in db.query(AppUser.role).all())
    if users_after != users_before or roles_after != roles_before:
        raise ScheduleReplacementError("Post-apply validation failed: users or roles changed")
    nearest = (
        db.query(PprNotification)
        .join(PprEvent)
        .order_by(PprNotification.scheduled_at.asc(), PprNotification.id.asc())
        .limit(10)
        .all()
    )
    missing_date_events = (
        db.query(PprEvent)
        .filter(PprEvent.date.is_(None))
        .order_by(PprEvent.id.asc())
        .limit(10)
        .all()
    )
    return {
        "preview": preview,
        "backup_path": str(backup_path),
        "deleted_ppr_events": preview["current_db"]["ppr_events"],
        "deleted_notifications": preview["current_db"]["ppr_notifications"],
        "created_ppr_events": preview["new_excel"]["unique_ppr"],
        "created_notifications": preview["new_excel"]["notifications_to_create"],
        "nearest_notifications": [
            {
                "notification_id": notif.id,
                "title": notif.event.title,
                "project": notif.event.project,
                "scheduled_at": notif.scheduled_at.isoformat(),
                "status": notif.status,
                "auto_send_enabled": notif.auto_send_enabled,
            }
            for notif in nearest
        ],
        "missing_date_events": [
            {
                "ppr_event_id": event.id,
                "title": event.title,
                "project": event.project,
                "source_key": event.source_key,
                "ppr_status": event.ppr_status,
            }
            for event in missing_date_events
        ],
    }
