import asyncio
import hashlib
import hmac
import json
import os
import tempfile
import time as time_module
import unittest
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from urllib.parse import urlencode

from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api.routes import router
from app.config import get_settings
from app.db.models import AppUser, AuditLog, PprEvent, PprNotification
from app.db.session import Base, get_db
from app.excel.import_service import (
    IMPORT_MODE_SAFE,
    apply_saved_import_preview,
    build_import_preview,
    save_import_preview,
)
from app.excel.importer import HEADER_ALIASES
from app.services.ppr_excel_export import (
    INFO_SHEET_NAME,
    MAIN_SHEET_NAME,
    SERVICE_SHEET_NAME,
    XLSX_MEDIA_TYPE,
    build_ppr_excel_export,
)
from app.services.statuses import PPR_STATUS_ARCHIVED, PPR_STATUS_IN_PROGRESS, PPR_STATUS_VERIFIED
from app.services.user_service import ROLE_ADMIN, ROLE_CHECKER


class PprExcelExportTestCase(unittest.TestCase):
    def setUp(self):
        handle = tempfile.NamedTemporaryFile(prefix="ppr-export-", suffix=".db", delete=False)
        handle.close()
        self.db_path = handle.name
        self.engine = create_engine(
            f"sqlite:///{self.db_path}",
            connect_args={"check_same_thread": False},
            future=True,
        )
        self.SessionLocal = sessionmaker(bind=self.engine, expire_on_commit=False, future=True)
        Base.metadata.create_all(self.engine)

        self.environment = {
            name: os.environ.get(name)
            for name in (
                "APP_ENV",
                "DEV_COMMANDS_ENABLED",
                "TELEGRAM_BOT_TOKEN",
                "TELEGRAM_WEBAPP_AUTH_MAX_AGE_SECONDS",
                "ADMIN_TELEGRAM_IDS",
                "OUTLOOK_ENABLED",
            )
        }
        os.environ.update(
            {
                "APP_ENV": "production",
                "DEV_COMMANDS_ENABLED": "false",
                "TELEGRAM_BOT_TOKEN": "unit-test-token",
                "TELEGRAM_WEBAPP_AUTH_MAX_AGE_SECONDS": "86400",
                "ADMIN_TELEGRAM_IDS": "",
                "OUTLOOK_ENABLED": "false",
            }
        )
        get_settings.cache_clear()

        self.tomorrow = date.today() + timedelta(days=10)
        with self.SessionLocal() as db:
            db.add_all(
                [
                    AppUser(
                        telegram_id="9001",
                        username="admin",
                        full_name="Admin User",
                        role=ROLE_ADMIN,
                        is_active=True,
                    ),
                    AppUser(
                        telegram_id="2001",
                        username="checker",
                        full_name="Checker User",
                        role=ROLE_CHECKER,
                        is_active=True,
                    ),
                    AppUser(
                        telegram_id="9002",
                        username="inactive",
                        full_name="Inactive Admin",
                        role=ROLE_ADMIN,
                        is_active=False,
                    ),
                ]
            )
            db.flush()

            imported = PprEvent(
                external_id="EXCEL-001",
                source_key="id:EXCEL-001",
                source_row=17,
                date=self.tomorrow,
                start_time=time(8, 15, 30),
                end_time=time(9, 45),
                notification_type="Плановое",
                project="Core",
                title="Обновление платформы",
                activities="Шаг 1\nШаг 2",
                responsible_setup="Инженер",
                responsible_report="Дежурный",
                source_link="https://example.test/ppr/1",
                notify_start=True,
                notify_end=False,
                is_active=True,
                ppr_status=PPR_STATUS_IN_PROGRESS,
                comment="Первая строка\nВторая строка",
            )
            manual = PprEvent(
                external_id="MANUAL-001",
                source_key=None,
                source_row=None,
                date=None,
                start_time=None,
                end_time=None,
                project="Manual",
                title="ППР без даты",
                activities="Дата будет назначена позже",
                notify_start=True,
                notify_end=False,
                is_active=True,
                is_manually_edited=True,
                comment="Создано вручную",
            )
            archived = PprEvent(
                external_id="AUTO-ARCHIVED",
                source_key="fingerprint:archived-event",
                source_row=33,
                date=self.tomorrow + timedelta(days=1),
                start_time=time(11, 0),
                project="Legacy",
                title="Архивная ППР",
                notify_start=True,
                notify_end=False,
                is_active=False,
                ppr_status=PPR_STATUS_ARCHIVED,
            )
            sent = PprEvent(
                external_id="EXCEL-SENT",
                source_key="id:EXCEL-SENT",
                source_row=None,
                date=self.tomorrow + timedelta(days=2),
                start_time=time(12, 0),
                project="Messaging",
                title="Уже отправленная ППР",
                notify_start=True,
                notify_end=False,
                is_active=True,
                ppr_status=PPR_STATUS_VERIFIED,
            )
            db.add_all([imported, manual, archived, sent])
            db.flush()
            db.add_all(
                [
                    PprNotification(
                        ppr_event_id=imported.id,
                        type="start",
                        scheduled_at=datetime.combine(imported.date, imported.start_time),
                        status="planned",
                        auto_send_enabled=True,
                        taken_by_id="2001",
                        taken_by_name="@checker",
                        taken_at=datetime.utcnow(),
                    ),
                    PprNotification(
                        ppr_event_id=archived.id,
                        type="start",
                        scheduled_at=datetime.combine(archived.date, archived.start_time),
                        status="cancelled",
                        auto_send_enabled=False,
                    ),
                    PprNotification(
                        ppr_event_id=sent.id,
                        type="start",
                        scheduled_at=datetime.combine(sent.date, sent.start_time),
                        status="sent",
                        auto_send_enabled=False,
                        telegram_chat_id="test-chat",
                        telegram_message_id="test-message",
                        sent_at=datetime.utcnow(),
                        checked_by_id="9001",
                        checked_by_name="@admin",
                        checked_at=datetime.utcnow(),
                    ),
                ]
            )
            db.commit()

        self.app = FastAPI()
        self.app.include_router(router)

        def override_db():
            db = self.SessionLocal()
            try:
                yield db
            finally:
                db.close()

        self.app.dependency_overrides[get_db] = override_db

    def tearDown(self):
        self.engine.dispose()
        for name, value in self.environment.items():
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value
        get_settings.cache_clear()
        try:
            os.remove(self.db_path)
        except FileNotFoundError:
            pass

    @staticmethod
    def _telegram_user(telegram_id: str, username: str, first_name: str, last_name: str) -> dict:
        return {
            "id": int(telegram_id),
            "username": username,
            "first_name": first_name,
            "last_name": last_name,
        }

    def init_data(self, telegram_id: str, username: str, first_name: str, last_name: str) -> str:
        values = {
            "auth_date": str(int(time_module.time())),
            "query_id": f"query-{telegram_id}",
            "user": json.dumps(
                self._telegram_user(telegram_id, username, first_name, last_name),
                ensure_ascii=False,
                separators=(",", ":"),
            ),
        }
        data_check_string = "\n".join(f"{key}={value}" for key, value in sorted(values.items()))
        secret_key = hmac.new(b"WebAppData", b"unit-test-token", hashlib.sha256).digest()
        values["hash"] = hmac.new(secret_key, data_check_string.encode(), hashlib.sha256).hexdigest()
        return urlencode(values)

    async def asgi_get(self, path: str, headers: dict[str, str] | None = None) -> tuple[int, dict[str, str], bytes]:
        request_pending = True
        sent_messages = []

        async def receive():
            nonlocal request_pending
            if request_pending:
                request_pending = False
                return {"type": "http.request", "body": b"", "more_body": False}
            return {"type": "http.disconnect"}

        async def send(message):
            sent_messages.append(message)

        raw_headers = [
            (name.lower().encode("latin-1"), value.encode("latin-1"))
            for name, value in (headers or {}).items()
        ]
        scope = {
            "type": "http",
            "asgi": {"version": "3.0", "spec_version": "2.3"},
            "http_version": "1.1",
            "method": "GET",
            "scheme": "http",
            "path": path,
            "raw_path": path.encode("ascii"),
            "query_string": b"",
            "root_path": "",
            "headers": raw_headers,
            "client": ("127.0.0.1", 12345),
            "server": ("testserver", 80),
        }
        await self.app(scope, receive, send)
        start = next(message for message in sent_messages if message["type"] == "http.response.start")
        response_headers = {
            name.decode("latin-1"): value.decode("latin-1")
            for name, value in start["headers"]
        }
        body = b"".join(
            message.get("body", b"")
            for message in sent_messages
            if message["type"] == "http.response.body"
        )
        return start["status"], response_headers, body

    @staticmethod
    def _table_snapshot(db, model) -> list[tuple]:
        columns = list(model.__table__.columns)
        return [
            tuple(getattr(item, column.name) for column in columns)
            for item in db.query(model).order_by(model.id.asc()).all()
        ]

    @staticmethod
    def _workbook_content(workbook) -> bytes:
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _row_for_id(worksheet, external_id: str) -> int:
        id_column = list(HEADER_ALIASES).index("ID") + 1
        for row_index in range(2, worksheet.max_row + 1):
            if worksheet.cell(row_index, id_column).value == external_id:
                return row_index
        raise AssertionError(f"ID not found: {external_id}")

    def test_admin_downloads_complete_import_compatible_workbook_and_audit_is_written(self):
        with self.SessionLocal() as db:
            before_events = self._table_snapshot(db, PprEvent)
            before_notifications = self._table_snapshot(db, PprNotification)
            before_audit_count = db.query(AuditLog).count()
        status, headers, body = asyncio.run(
            self.asgi_get(
                "/api/export/ppr.xlsx",
                {"X-Telegram-Init-Data": self.init_data("9001", "admin", "Admin", "User")},
            )
        )

        self.assertEqual(status, 200)
        self.assertEqual(headers["content-type"], XLSX_MEDIA_TYPE)
        self.assertRegex(
            headers["content-disposition"],
            r'^attachment; filename="PPR_export_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}\.xlsx"$',
        )
        workbook = load_workbook(BytesIO(body), data_only=False)
        self.assertEqual(workbook.sheetnames, [MAIN_SHEET_NAME, SERVICE_SHEET_NAME, INFO_SHEET_NAME])

        worksheet = workbook[MAIN_SHEET_NAME]
        self.assertEqual([cell.value for cell in worksheet[1]], list(HEADER_ALIASES))
        self.assertEqual(worksheet.max_row, 5)
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertTrue(worksheet.auto_filter.ref)

        exported_ids = {worksheet.cell(row, 1).value for row in range(2, worksheet.max_row + 1)}
        self.assertEqual(exported_ids, {"EXCEL-001", "MANUAL-001", "AUTO-ARCHIVED", "EXCEL-SENT"})
        imported_row = self._row_for_id(worksheet, "EXCEL-001")
        undated_row = self._row_for_id(worksheet, "MANUAL-001")
        date_column = list(HEADER_ALIASES).index("Дата") + 1
        start_column = list(HEADER_ALIASES).index("Время выхода") + 1
        activities_column = list(HEADER_ALIASES).index("Активности") + 1
        comment_column = list(HEADER_ALIASES).index("Комментарий") + 1
        self.assertEqual(worksheet.cell(imported_row, date_column).value.date(), self.tomorrow)
        self.assertEqual(worksheet.cell(imported_row, start_column).value, time(8, 15, 30))
        self.assertEqual(worksheet.cell(imported_row, activities_column).value, "Шаг 1\nШаг 2")
        self.assertEqual(worksheet.cell(imported_row, comment_column).value, "Первая строка\nВторая строка")
        self.assertIsNone(worksheet.cell(undated_row, date_column).value)
        self.assertIsNone(worksheet.cell(undated_row, start_column).value)
        self.assertTrue(all(cell.data_type != "f" for row in worksheet.iter_rows() for cell in row))

        with self.SessionLocal() as db:
            self.assertEqual(self._table_snapshot(db, PprEvent), before_events)
            self.assertEqual(self._table_snapshot(db, PprNotification), before_notifications)
            self.assertEqual(db.query(AuditLog).count(), before_audit_count + 1)
            audit = db.query(AuditLog).filter(AuditLog.action == "ppr_excel_export").one()
            self.assertEqual(audit.user_id, "9001")
            self.assertIn("app_user_id=", audit.comment)
            self.assertIn("telegram_id=9001", audit.comment)
            self.assertIn("exported_count=4", audit.comment)
            self.assertIn("filename=PPR_export_", audit.comment)
            self.assertNotIn("query-9001", audit.comment)

    def test_checker_inactive_invalid_and_missing_auth_are_rejected_without_audit(self):
        checker, _, _ = asyncio.run(
            self.asgi_get(
                "/api/export/ppr.xlsx",
                {"X-Telegram-Init-Data": self.init_data("2001", "checker", "Checker", "User")},
            )
        )
        inactive, _, _ = asyncio.run(
            self.asgi_get(
                "/api/export/ppr.xlsx",
                {"X-Telegram-Init-Data": self.init_data("9002", "inactive", "Inactive", "Admin")},
            )
        )
        invalid, _, _ = asyncio.run(
            self.asgi_get(
                "/api/export/ppr.xlsx",
                {"X-Telegram-Init-Data": "not-valid-init-data"},
            )
        )
        missing, _, _ = asyncio.run(self.asgi_get("/api/export/ppr.xlsx"))

        self.assertEqual(checker, 403)
        self.assertEqual(inactive, 403)
        self.assertEqual(invalid, 401)
        self.assertEqual(missing, 401)
        with self.SessionLocal() as db:
            self.assertEqual(db.query(AuditLog).count(), 0)

    def test_export_route_is_registered_as_get(self):
        route = next(item for item in router.routes if item.path == "/api/export/ppr.xlsx")
        self.assertEqual(route.methods, {"GET"})

    def test_unchanged_export_preview_and_apply_do_not_mutate_events_or_notifications(self):
        with self.SessionLocal() as db:
            admin = db.query(AppUser).filter(AppUser.telegram_id == "9001").one()
            export = build_ppr_excel_export(db, admin)
            preview = build_import_preview(db, export.content, export.filename, IMPORT_MODE_SAFE)

            self.assertEqual(preview["summary"]["new_events"], 0)
            self.assertEqual(preview["summary"]["updated_events"], 0)
            self.assertEqual(preview["summary"]["duplicate_rows"], 0)
            self.assertEqual(preview["summary"]["errors_count"], 0)
            self.assertEqual(preview["summary"]["skipped_manual_events"], 1)
            self.assertTrue(
                all(
                    item["notification_action"] == "skip"
                    for item in preview["details"]
                    if item.get("excel_row_number")
                )
            )

            before_events = self._table_snapshot(db, PprEvent)
            before_notifications = self._table_snapshot(db, PprNotification)
            saved = save_import_preview(db, export.content, export.filename, IMPORT_MODE_SAFE, admin)
            result = asyncio.run(
                apply_saved_import_preview(
                    db,
                    saved["preview_id"],
                    export.content,
                    IMPORT_MODE_SAFE,
                    admin,
                )
            )

            self.assertEqual(result["applied"]["created"], 0)
            self.assertEqual(result["applied"]["updated"], 0)
            self.assertEqual(result["applied"]["notifications_synced"], 0)
            self.assertEqual(self._table_snapshot(db, PprEvent), before_events)
            self.assertEqual(self._table_snapshot(db, PprNotification), before_notifications)

    def test_changed_export_matches_existing_id_and_reschedules_only_planned_notification(self):
        with self.SessionLocal() as db:
            admin = db.query(AppUser).filter(AppUser.telegram_id == "9001").one()
            export = build_ppr_excel_export(db, admin)
            workbook = load_workbook(BytesIO(export.content))
            worksheet = workbook[MAIN_SHEET_NAME]
            row_index = self._row_for_id(worksheet, "EXCEL-001")
            title_column = list(HEADER_ALIASES).index("Название ППР") + 1
            date_column = list(HEADER_ALIASES).index("Дата") + 1
            time_column = list(HEADER_ALIASES).index("Время выхода") + 1
            changed_date = self.tomorrow + timedelta(days=5)
            changed_time = time(10, 20, 30)
            worksheet.cell(row_index, title_column).value = "Обновленное название"
            worksheet.cell(row_index, date_column).value = changed_date
            worksheet.cell(row_index, time_column).value = changed_time
            changed_content = self._workbook_content(workbook)

            event = db.query(PprEvent).filter(PprEvent.external_id == "EXCEL-001").one()
            notification = db.query(PprNotification).filter(PprNotification.ppr_event_id == event.id).one()
            original_taken_by = notification.taken_by_id
            preview = build_import_preview(db, changed_content, export.filename, IMPORT_MODE_SAFE)
            detail = next(item for item in preview["details"] if item.get("external_id") == "EXCEL-001")
            self.assertEqual(detail["action"], "update")
            self.assertEqual(detail["ppr_event_id"], event.id)
            self.assertEqual(preview["summary"]["new_events"], 0)

            saved = save_import_preview(db, changed_content, export.filename, IMPORT_MODE_SAFE, admin)
            result = asyncio.run(
                apply_saved_import_preview(
                    db,
                    saved["preview_id"],
                    changed_content,
                    IMPORT_MODE_SAFE,
                    admin,
                )
            )
            db.refresh(event)
            db.refresh(notification)
            self.assertEqual(result["applied"]["updated"], 1)
            self.assertEqual(event.title, "Обновленное название")
            self.assertEqual(event.date, changed_date)
            self.assertEqual(event.start_time, changed_time)
            self.assertEqual(notification.scheduled_at, datetime.combine(changed_date, changed_time))
            self.assertEqual(notification.status, "planned")
            self.assertEqual(notification.taken_by_id, original_taken_by)

    def test_changed_sent_event_is_not_automatically_requeued(self):
        with self.SessionLocal() as db:
            admin = db.query(AppUser).filter(AppUser.telegram_id == "9001").one()
            export = build_ppr_excel_export(db, admin)
            workbook = load_workbook(BytesIO(export.content))
            worksheet = workbook[MAIN_SHEET_NAME]
            row_index = self._row_for_id(worksheet, "EXCEL-SENT")
            date_column = list(HEADER_ALIASES).index("Дата") + 1
            worksheet.cell(row_index, date_column).value = self.tomorrow + timedelta(days=20)
            changed_content = self._workbook_content(workbook)

            event = db.query(PprEvent).filter(PprEvent.external_id == "EXCEL-SENT").one()
            notification = db.query(PprNotification).filter(PprNotification.ppr_event_id == event.id).one()
            original_notification = self._table_snapshot(db, PprNotification)[-1]
            preview = build_import_preview(db, changed_content, export.filename, IMPORT_MODE_SAFE)
            detail = next(item for item in preview["details"] if item.get("external_id") == "EXCEL-SENT")
            self.assertEqual(detail["action"], "update")
            self.assertEqual(detail["notification_action"], "skip")
            self.assertIn("sent", detail["notification_reason"])

            saved = save_import_preview(db, changed_content, export.filename, IMPORT_MODE_SAFE, admin)
            result = asyncio.run(
                apply_saved_import_preview(
                    db,
                    saved["preview_id"],
                    changed_content,
                    IMPORT_MODE_SAFE,
                    admin,
                )
            )
            db.refresh(notification)
            self.assertEqual(result["applied"]["notifications_synced"], 0)
            self.assertEqual(self._table_snapshot(db, PprNotification)[-1], original_notification)
            self.assertEqual(notification.status, "sent")

    def test_text_change_does_not_reenable_cancelled_notification(self):
        with self.SessionLocal() as db:
            admin = db.query(AppUser).filter(AppUser.telegram_id == "9001").one()
            export = build_ppr_excel_export(db, admin)
            workbook = load_workbook(BytesIO(export.content))
            worksheet = workbook[MAIN_SHEET_NAME]
            row_index = self._row_for_id(worksheet, "AUTO-ARCHIVED")
            title_column = list(HEADER_ALIASES).index("Название ППР") + 1
            worksheet.cell(row_index, title_column).value = "Уточненная архивная ППР"
            changed_content = self._workbook_content(workbook)

            event = db.query(PprEvent).filter(PprEvent.external_id == "AUTO-ARCHIVED").one()
            notification = db.query(PprNotification).filter(PprNotification.ppr_event_id == event.id).one()
            original_notification = tuple(
                getattr(notification, column.name) for column in PprNotification.__table__.columns
            )
            preview = build_import_preview(db, changed_content, export.filename, IMPORT_MODE_SAFE)
            detail = next(item for item in preview["details"] if item.get("external_id") == "AUTO-ARCHIVED")
            self.assertEqual(detail["action"], "update")
            self.assertEqual(detail["notification_action"], "skip")

            saved = save_import_preview(db, changed_content, export.filename, IMPORT_MODE_SAFE, admin)
            asyncio.run(
                apply_saved_import_preview(
                    db,
                    saved["preview_id"],
                    changed_content,
                    IMPORT_MODE_SAFE,
                    admin,
                )
            )
            db.refresh(notification)
            current_notification = tuple(
                getattr(notification, column.name) for column in PprNotification.__table__.columns
            )
            self.assertEqual(current_notification, original_notification)
            self.assertEqual(notification.status, "cancelled")

    def test_filename_uses_configured_timezone(self):
        with self.SessionLocal() as db:
            admin = db.query(AppUser).filter(AppUser.telegram_id == "9001").one()
            export = build_ppr_excel_export(
                db,
                admin,
                now=datetime(2026, 1, 2, 0, 15, tzinfo=timezone.utc),
            )

        self.assertEqual(export.filename, "PPR_export_2026-01-02_03-15.xlsx")


if __name__ == "__main__":
    unittest.main()
