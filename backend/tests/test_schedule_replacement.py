import os
import tempfile
import unittest
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.config import get_settings
from app.db.models import AppUser, AuditLog, PprEvent, PprNotification
from app.db.session import Base
from app.excel.import_service import IMPORT_MODE_SAFE, build_import_preview
from app.services.schedule_replacement_service import (
    REPLACE_CONFIRMATION,
    ScheduleReplacementError,
    ScheduleValidationError,
    build_schedule_replacement_preview,
    export_validation_report,
    replace_schedule_from_excel,
)


HEADERS = [
    "ID", "Дата", "Время выхода", "Время завершения", "Тип уведомления", "Проект", "Название ППР",
    "Активности", "Ответственный настройка", "Ответственный отчетка", "Ссылка", "Уведомлять о выходе",
    "Уведомлять о завершении", "Активно", "Комментарий", "Исходная строка",
]


def workbook_bytes(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ППР_для_бота"
    worksheet.append(HEADERS)
    for item in rows:
        worksheet.append([
            item.get("id"), item.get("date"), item.get("start_time", time(9, 0)), None, None,
            item.get("project", "Project"), item.get("title", "PPR"), item.get("activities", "Work"),
            None, None, None, item.get("notify", True), False, item.get("active", True), None, item.get("source_row"),
        ])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class ScheduleReplacementTestCase(unittest.TestCase):
    def setUp(self):
        self.previous_auto_send = os.environ.get("NOTIFICATIONS_AUTO_SEND_ENABLED")
        os.environ["NOTIFICATIONS_AUTO_SEND_ENABLED"] = "false"
        get_settings.cache_clear()
        handle = tempfile.NamedTemporaryFile(prefix="ppr-replace-", suffix=".db", delete=False)
        handle.close()
        self.path = handle.name
        self.engine = create_engine(f"sqlite:///{self.path}", future=True)
        self.SessionLocal = sessionmaker(bind=self.engine, expire_on_commit=False, future=True)
        Base.metadata.create_all(self.engine)
        backup_fd, backup_path = tempfile.mkstemp(prefix="ppr-replace-", suffix=".dump")
        os.close(backup_fd)
        self.backup = Path(backup_path)
        self.backup.write_bytes(b"backup")
        self.future = date.today() + timedelta(days=10)

    def tearDown(self):
        self.engine.dispose()
        os.unlink(self.path)
        self.backup.unlink(missing_ok=True)
        if self.previous_auto_send is None:
            os.environ.pop("NOTIFICATIONS_AUTO_SEND_ENABLED", None)
        else:
            os.environ["NOTIFICATIONS_AUTO_SEND_ENABLED"] = self.previous_auto_send
        get_settings.cache_clear()

    def valid_content(self, *rows):
        return workbook_bytes(rows or ({"id": "NEW-1", "date": self.future, "title": "Future PPR"},))

    def add_old_schedule(self, db):
        user = AppUser(telegram_id="1", username="admin", role="admin", is_active=True)
        event = PprEvent(external_id="OLD-1", source_key="id:OLD-1", title="Old", date=self.future, start_time=time(8, 0))
        db.add_all([user, event])
        db.flush()
        notification = PprNotification(ppr_event_id=event.id, type="start", scheduled_at=datetime.combine(self.future, time(8, 0)), status="sent")
        db.add(notification)
        db.flush()
        db.add(AuditLog(ppr_event_id=event.id, notification_id=notification.id, action="old_action", user_id="1", user_name="admin"))
        db.commit()
        return user, event, notification

    def apply(self, db, content):
        return replace_schedule_from_excel(db, content, "schedule.xlsx", confirmation=REPLACE_CONFIRMATION, backup_creator=lambda: self.backup)

    def test_preview_never_changes_database(self):
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            preview = build_schedule_replacement_preview(db, self.valid_content(), "schedule.xlsx")
            self.assertFalse(preview["errors"])
            self.assertEqual(db.query(PprEvent).count(), 1)
            self.assertEqual(db.query(PprNotification).count(), 1)

    def test_apply_requires_exact_confirmation(self):
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            with self.assertRaises(ScheduleReplacementError):
                replace_schedule_from_excel(db, self.valid_content(), "schedule.xlsx", confirmation="wrong", backup_creator=lambda: self.backup)
            self.assertEqual(db.query(PprEvent).count(), 1)

    def test_invalid_excel_and_duplicates_do_not_delete_old_schedule(self):
        broken = workbook_bytes([
            {"id": "DUP", "date": self.future, "title": "A"},
            {"id": "DUP", "date": self.future, "title": "B"},
            {"id": "NO-DATE", "date": None, "start_time": None, "title": "No date"},
        ])
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            preview = build_schedule_replacement_preview(db, broken, "schedule.xlsx")
            self.assertTrue(preview["errors"])
            with self.assertRaises(ScheduleValidationError):
                self.apply(db, broken)
            self.assertEqual(db.query(PprEvent).count(), 1)
            self.assertEqual(db.query(PprNotification).count(), 1)

    def test_past_row_blocks_replacement(self):
        with self.SessionLocal() as db:
            preview = build_schedule_replacement_preview(db, workbook_bytes([{"id": "OLD", "date": date.today() - timedelta(days=1)}]), "schedule.xlsx")
            self.assertTrue(any("в будущем" in item["reason"] for item in preview["errors"]))

    def test_missing_date_and_time_creates_ppr_without_notification(self):
        content = workbook_bytes([{"date": None, "start_time": None, "title": "Needs date"}])
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            preview = build_schedule_replacement_preview(db, content, "schedule.xlsx")
            self.assertFalse(preview["errors"])
            self.assertEqual(preview["new_excel"]["missing_date_ppr"], 1)
            self.assertEqual(preview["new_excel"]["notifications_to_create"], 0)
            self.apply(db, content)
            event = db.query(PprEvent).one()
            self.assertIsNone(event.date)
            self.assertIsNone(event.start_time)
            self.assertEqual(event.ppr_status, "scheduled")
            self.assertTrue(event.source_key.startswith("replace-row:ППР_для_бота:"))
            self.assertEqual(db.query(PprNotification).count(), 0)

    def test_date_without_time_is_blocking_error(self):
        content = workbook_bytes([{"id": "TIME-1", "date": self.future, "start_time": None, "title": "No time"}])
        with self.SessionLocal() as db:
            preview = build_schedule_replacement_preview(db, content, "schedule.xlsx")
            self.assertTrue(any("Время выхода не заполнено" in item["reason"] for item in preview["errors"]))

    def test_equal_fingerprints_are_warnings_and_get_unique_replace_keys(self):
        content = workbook_bytes([
            {"date": None, "start_time": None, "title": "Repeated"},
            {"date": None, "start_time": None, "title": "Repeated"},
        ])
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            preview = build_schedule_replacement_preview(db, content, "schedule.xlsx")
            self.assertFalse(preview["errors"])
            self.assertEqual(preview["new_excel"]["ambiguous_fingerprint_groups"], 1)
            self.assertEqual(preview["new_excel"]["warnings"], 1)
            self.apply(db, content)
            self.assertEqual(db.query(PprEvent).count(), 2)
            self.assertEqual(db.query(PprEvent.source_key).distinct().count(), 2)

    def test_incremental_import_keeps_its_existing_fingerprint_behavior(self):
        content = workbook_bytes([
            {"date": None, "start_time": None, "title": "Repeated"},
            {"date": None, "start_time": None, "title": "Repeated"},
        ])
        with self.SessionLocal() as db:
            preview = build_import_preview(db, content, "schedule.xlsx", IMPORT_MODE_SAFE)
            self.assertEqual(preview["summary"]["duplicate_rows"], 2)

    def test_backup_failure_blocks_replacement(self):
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            with self.assertRaises(ScheduleReplacementError):
                replace_schedule_from_excel(db, self.valid_content(), "schedule.xlsx", confirmation=REPLACE_CONFIRMATION, backup_creator=lambda: Path("missing.dump"))
            self.assertEqual(db.query(PprEvent).count(), 1)

    def test_apply_replaces_schedule_preserves_users_and_audit(self):
        content = self.valid_content(
            {"id": "NEW-1", "date": self.future, "title": "Future one"},
            {"id": "NEW-2", "date": self.future + timedelta(days=1), "title": "Future two", "notify": False},
        )
        with self.SessionLocal() as db:
            _, old_event, old_notification = self.add_old_schedule(db)
            result = self.apply(db, content)
            self.assertEqual(result["deleted_ppr_events"], 1)
            self.assertEqual(result["deleted_notifications"], 1)
            self.assertEqual(db.query(AppUser).count(), 1)
            self.assertEqual(db.query(PprEvent).count(), 2)
            self.assertEqual(db.query(PprNotification).count(), 1)
            self.assertIsNone(db.query(AuditLog).filter(AuditLog.action == "old_action").one().ppr_event_id)
            self.assertEqual(db.query(AuditLog).filter(AuditLog.action == "old_action").one().notification_id, None)
            summary = db.query(AuditLog).filter(AuditLog.action == "schedule_replaced_from_excel").one()
            self.assertIn("created_ppr=2", summary.comment)
            self.assertEqual(db.query(PprEvent).filter(PprEvent.external_id == "OLD-1").count(), 0)

    def test_rollback_keeps_old_schedule_when_event_creation_fails(self):
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            with patch("app.services.schedule_replacement_service._create_events_from_rows", side_effect=RuntimeError("boom")):
                with self.assertRaises(RuntimeError):
                    self.apply(db, self.valid_content())
            self.assertEqual(db.query(PprEvent).count(), 1)
            self.assertEqual(db.query(PprNotification).count(), 1)

    def test_repeat_replacement_creates_a_clean_set_without_duplicates(self):
        content = self.valid_content({"id": "NEW-1", "date": self.future, "title": "Future"})
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            self.apply(db, content)
            self.apply(db, content)
            self.assertEqual(db.query(PprEvent).count(), 1)
            self.assertEqual(db.query(PprNotification).count(), 1)
            self.assertEqual(db.query(PprEvent.source_key).distinct().count(), 1)

    def test_auto_send_and_processing_block_replacement(self):
        with self.SessionLocal() as db:
            _, _, notification = self.add_old_schedule(db)
            notification.status = "processing"
            db.commit()
            with self.assertRaises(ScheduleReplacementError):
                self.apply(db, self.valid_content())
            self.assertEqual(db.query(PprEvent).count(), 1)

    def test_validation_report_has_expected_sheets_and_does_not_change_database(self):
        content = self.valid_content()
        with self.SessionLocal() as db:
            self.add_old_schedule(db)
            before = (db.query(PprEvent).count(), db.query(PprNotification).count(), db.query(AppUser).count())
            preview = build_schedule_replacement_preview(db, content, "schedule.xlsx")
            handle = tempfile.NamedTemporaryFile(prefix="ppr-validation-", suffix=".xlsx", delete=False)
            handle.close()
            report_path = Path(handle.name)
            try:
                export_validation_report(content, report_path, preview)
                workbook = load_workbook(report_path, data_only=True)
                self.assertEqual(workbook.sheetnames, ["Summary", "Errors", "DuplicateGroups", "Без даты", "ValidRows"])
                for sheet in workbook.worksheets:
                    self.assertEqual(sheet.freeze_panes, "A2")
                    self.assertTrue(sheet.auto_filter.ref)
                self.assertEqual(before, (db.query(PprEvent).count(), db.query(PprNotification).count(), db.query(AppUser).count()))
            finally:
                report_path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
